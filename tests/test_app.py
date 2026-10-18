import importlib.util
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook

APP_PATH = Path(__file__).resolve().parents[1] / 'app.py'
spec = importlib.util.spec_from_file_location('educacao_app', APP_PATH)
app = importlib.util.module_from_spec(spec)
spec.loader.exec_module(app)

class FinanceTests(unittest.TestCase):
    def test_parse_amount_brazilian(self):
        self.assertAlmostEqual(app.parse_amount('1.234,56'), 1234.56)
        self.assertAlmostEqual(app.parse_amount('125,50'), 125.50)
        self.assertAlmostEqual(app.parse_amount(42), 42.0)

    def test_workbook_has_expected_sections(self):
        self.assertEqual(set(app.GROUPS), {'Essenciais','Metas','Gastos livres','Educação','Investimentos'})
        self.assertIn(111, app.GROUPS['Investimentos'])
        self.assertIn(21, app.GROUPS['Essenciais'])

    def test_target_percentages_sum_to_100(self):
        self.assertAlmostEqual(sum(app.TARGETS_DEFAULT.values()),100.0)

    def test_import_reads_monthly_values_without_totals(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'financeiro.xlsx'
            wb=Workbook(); ws=wb.active; ws.title='Orçamento Mensal'
            ws['D2']=2026; ws['B13']=0.55; ws['B14']=0.20; ws['B15']=0.10; ws['B16']=0.10; ws['B17']=0.05
            ws['C5']='Receita Marido'; ws['D5']=3200
            ws['C21']='Prestação Carro'; ws['D21']=450
            ws['C111']='Ações'; ws['E111']=100
            ws['C82']='TOTAL ESSENCIAL'; ws['D82']=450
            wb.save(path)
            year,entries,opening,targets=app.extract_workbook_data(path)
            self.assertEqual(year,2026)
            self.assertEqual(len(entries),3)
            self.assertEqual(entries[0][1:4],('Receita','Receita Marido',3200.0))
            self.assertEqual(entries[1][4],'Essenciais')
            self.assertEqual(entries[2][1],'Investimento')
            self.assertAlmostEqual(targets['Essenciais'],55.0)
            self.assertEqual(opening,0.0)

    def test_short_money_and_axis(self):
        self.assertEqual(app.short_money(950),'950')
        self.assertEqual(app.short_money(1000),'1 mil')
        self.assertEqual(app.short_money(2500),'2,5 mil')
        self.assertEqual(app.short_money(2_000_000),'2 mi')
        self.assertEqual(app.nice_max(0),1.0)
        self.assertEqual(app.nice_max(4200),4800)
        self.assertEqual(app.nice_max(24000),24000)
        self.assertEqual(app.nice_max(24100),32000)
        for v in (1,37,950,4200,24000,123456): self.assertGreaterEqual(app.nice_max(v),v)

    def test_top_n_groups_remaining_in_outras(self):
        items=[('a',5),('b',3),('c',2),('d',1),('e',0)]
        self.assertEqual(app.top_n(items,3),[('a',5),('b',3),('Outras',3)])
        self.assertEqual(app.top_n(items[:2],3),[('a',5),('b',3)])

    def test_colorize_keeps_fixed_group_colors(self):
        out=app.colorize([('Essenciais',1),('Outro',2)],app.GROUP_COLORS)
        self.assertEqual(out[0][2],app.GROUP_COLORS['Essenciais'])
        self.assertEqual(len(out[1][2]),7)

    def test_month_label_and_safe_amount(self):
        self.assertEqual(app.month_label('2026-10'),'Outubro de 2026')
        self.assertEqual(app.safe_amount('1.234,56'),1234.56)
        self.assertEqual(app.safe_amount('abc'),0.0)
        self.assertEqual(app.safe_amount(''),0.0)

    def test_months_until(self):
        self.assertIsNone(app.months_until(''))
        self.assertIsNone(app.months_until('lixo'))
        self.assertIsInstance(app.months_until('2099-01'),int)

    def test_every_group_has_a_color(self):
        self.assertEqual(set(app.GROUP_COLORS),set(app.TARGETS_DEFAULT))

    # ---------- estrutura da planilha e simuladores
    def test_structure_matches_workbook_rows(self):
        for group,rows in app.GROUPS.items():
            n=sum(len(cats) for g,_,cats in app.STRUCTURE if g==group)
            self.assertEqual(n,len(rows),group)
        self.assertEqual(sum(len(c) for _,_,c in app.STRUCTURE),74)
        self.assertEqual(len([1 for g,s,_ in app.STRUCTURE if g=='Essenciais' and s]),10)

    def test_split_income_and_reserve(self):
        parts=app.split_income(5000,app.TARGETS_DEFAULT)
        self.assertAlmostEqual(sum(v for _,_,v in parts),5000)
        self.assertAlmostEqual(dict((g,v) for g,_,v in parts)['Essenciais'],2750)
        self.assertEqual(app.reserve_target(3000,6),18000)
        self.assertEqual(app.months_to_goal(18000,1500),12)
        self.assertIsNone(app.months_to_goal(100,0))
        self.assertEqual(app.duration_text(14),'1 ano e 2 meses')
        self.assertEqual(app.duration_text(5),'5 meses')

    def test_compound_projection(self):
        self.assertAlmostEqual(app.compound_projection(100,0,2)[-1][2],2400)
        self.assertAlmostEqual(app.compound_projection(0,12,1,1000)[0][2],1120,places=6)
        self.assertEqual(len(app.compound_projection(100,8,10)),10)

    # ---------- cálculo do ano (equivale às fórmulas da planilha)
    def make_model(self):
        tmp=tempfile.mkdtemp(); app.DB=Path(tmp)/'t.db'
        names=['month_rows','sums','by_category','opening_balance','targets','pref','year_model','block_series','grid_rows','set_cell','store_entry','store_goal','scope_filter','clear_scope','wipe_all','safety_backup','get_month','available_months']
        H=type('H',(),{n:app.App.__dict__[n] for n in names}); h=H(); h.con=app.connect(); h.top_month=type('V',(),{'v':'2026-05','get':lambda s:s.v,'set':lambda s,x:setattr(s,'v',x)})(); return h

    def add(self,h,data,tipo,cat,valor,origem='manual',grupo=''):
        h.con.execute('INSERT INTO lancamentos(data,tipo,categoria,descricao,valor,origem,grupo) VALUES(?,?,?,?,?,?,?)',(data,tipo,cat,'',valor,origem,grupo)); h.con.commit()

    def test_carry_over_only_positive_balance(self):
        h=self.make_model()
        self.add(h,'2026-01-05','Receita','Receita Marido',1000); self.add(h,'2026-01-06','Despesa','Mercado',1500,grupo='Essenciais')
        self.add(h,'2026-02-05','Receita','Receita Marido',100); self.add(h,'2026-03-05','Receita','Receita Marido',50)
        m=h.year_model(2026)
        self.assertEqual(m['saldo'][0],-500)
        self.assertEqual(m['opening'][1],0)        # saldo negativo não é levado adiante
        self.assertEqual(m['saldo'][1],100)
        self.assertEqual(m['opening'][2],100)      # saldo positivo vira a sobra do mês seguinte
        self.assertAlmostEqual(h.opening_balance('2026-03'),100)

    def test_opening_anchor_and_totals(self):
        h=self.make_model(); h.con.execute("INSERT INTO saldos_iniciais VALUES('2026-01',200)"); h.con.commit()
        self.add(h,'2026-01-05','Receita','Receita Marido',1000)
        self.add(h,'2026-01-06','Despesa','Prestação Carro',300,grupo='Essenciais')
        self.add(h,'2026-01-07','Investimento','Ações',100,grupo='Investimentos')
        m=h.year_model(2026)
        self.assertEqual(m['opening'][0],200); self.assertEqual(m['despesas'][0],300); self.assertEqual(m['saidas'][0],400)
        self.assertEqual(m['saldo'][0],800); self.assertEqual(m['patrimonio'][11],100)
        self.assertAlmostEqual(m['orcado']['Essenciais'][0],550)

    def test_ambiguous_outros_and_unclassified(self):
        h=self.make_model()
        self.add(h,'2026-05-05','Despesa','Outros',10,grupo='Gastos livres'); self.add(h,'2026-05-05','Despesa','Outros',20,grupo='Educação')
        self.add(h,'2026-05-05','Investimento','Outros',30,grupo='Investimentos'); self.add(h,'2026-05-06','Despesa','Categoria antiga',7)
        m=h.year_model(2026)
        self.assertEqual(m['cell'][('Gastos livres','Outros')][4],10); self.assertEqual(m['cell'][('Educação','Outros')][4],20)
        self.assertEqual(m['cell'][('Investimentos','Outros')][4],30); self.assertEqual(m['cell'][('Essenciais',app.OTHER_ROW)][4],7)
        self.assertTrue(any(app.OTHER_ROW in r[1] for r in h.grid_rows(m,2026)))

    def test_grid_cell_edit_respects_manual_entries(self):
        h=self.make_model()
        self.add(h,'2026-04-01','Despesa','Gasolina',35,origem='planilha',grupo='Essenciais')
        self.assertIsNone(h.set_cell(2026,4,'Essenciais','Gasolina',777))
        self.assertEqual(h.year_model(2026)['cell'][('Essenciais','Gasolina')][3],777)
        self.add(h,'2026-04-10','Despesa','Gasolina',100,grupo='Essenciais')
        self.assertIsNone(h.set_cell(2026,4,'Essenciais','Gasolina',900))
        self.assertEqual(h.year_model(2026)['cell'][('Essenciais','Gasolina')][3],900)
        self.assertIsNotNone(h.set_cell(2026,4,'Essenciais','Gasolina',50))   # menor que o lançamento manual
        self.assertEqual(h.year_model(2026)['cell'][('Essenciais','Gasolina')][3],900)

    def test_grid_rows_shape(self):
        h=self.make_model(); rows=h.grid_rows(h.year_model(2026),2026)
        self.assertTrue(all(len(r[2])==12 for r in rows)); self.assertEqual(rows[0][1],'1. RECEITAS')
        self.assertEqual(len({r[0] for r in rows}),len(rows))

    # ---------- corrigir, limpar e restaurar
    def test_edit_entry_rewrites_and_becomes_manual(self):
        h=self.make_model(); self.add(h,'2026-05-02','Despesa','Mercado',100,origem='planilha',grupo='Essenciais')
        i=h.con.execute('SELECT id FROM lancamentos').fetchone()[0]
        h.store_entry('2026-05-09','Despesa','Padaria','corrigido',42.5,'Essenciais',i)
        r=h.con.execute('SELECT data,categoria,descricao,valor,origem FROM lancamentos WHERE id=?',(i,)).fetchone()
        self.assertEqual(r,('2026-05-09','Padaria','corrigido',42.5,'manual'))
        self.assertEqual(h.con.execute('SELECT COUNT(*) FROM lancamentos').fetchone()[0],1)   # reescreveu, não duplicou

    def test_edit_goal_rewrites(self):
        h=self.make_model(); h.store_goal('Viagem',6000,1500,'2027-06')
        i=h.con.execute('SELECT id FROM metas').fetchone()[0]; h.store_goal('Viagem Europa',9000,2000,'',i)
        self.assertEqual(h.con.execute('SELECT nome,valor_alvo,valor_atual,prazo FROM metas').fetchall(),[('Viagem Europa',9000,2000,'')])

    def test_clear_scopes(self):
        h=self.make_model(); h.con.execute("INSERT INTO saldos_iniciais VALUES('2026-01',50)")
        self.add(h,'2026-05-02','Despesa','Mercado',10,grupo='Essenciais'); self.add(h,'2026-06-02','Despesa','Mercado',20,grupo='Essenciais')
        self.add(h,'2026-06-03','Despesa','Gás',30,origem='planilha',grupo='Essenciais'); self.add(h,'2025-06-03','Despesa','Gás',40,grupo='Essenciais')
        self.assertEqual(h.clear_scope('month'),1)                 # só maio/2026
        self.assertEqual(h.clear_scope('imported'),1)              # só o que veio da planilha
        self.assertEqual(h.con.execute('SELECT COUNT(*) FROM lancamentos').fetchone()[0],2)
        self.assertEqual(h.clear_scope('year'),1)                  # 2026: sobra o de 2025
        self.assertEqual(h.con.execute('SELECT valor FROM lancamentos').fetchall(),[(40,)])
        self.assertEqual(h.con.execute('SELECT COUNT(*) FROM saldos_iniciais').fetchone()[0],0)
        with self.assertRaises(ValueError): h.scope_filter('xyz')

    def test_wipe_all_resets_defaults_and_backup_keeps_old_data(self):
        h=self.make_model(); self.add(h,'2026-05-02','Despesa','Mercado',10,grupo='Essenciais'); h.store_goal('M',10,0,'')
        h.con.execute("UPDATE preferencias SET valor='70' WHERE chave='alvo_Essenciais'"); h.con.commit()
        backup=h.safety_backup('teste'); h.wipe_all()
        self.assertEqual(h.con.execute('SELECT COUNT(*) FROM lancamentos').fetchone()[0],0)
        self.assertEqual(h.con.execute('SELECT COUNT(*) FROM metas').fetchone()[0],0)
        self.assertEqual(h.targets(),app.TARGETS_DEFAULT)
        import sqlite3
        old=sqlite3.connect(backup); self.assertEqual(old.execute('SELECT COUNT(*) FROM lancamentos').fetchone()[0],1); old.close()

    # ---------- anos sem limite
    def test_carry_over_has_no_year_limit(self):
        h=self.make_model(); self.add(h,'2026-01-05','Receita','Receita Marido',100)
        self.assertEqual(h.opening_balance('2026-01'),0)
        self.assertEqual(h.opening_balance('2026-02'),100)
        self.assertEqual(h.opening_balance('2031-06'),100)      # 5 anos depois
        self.assertEqual(h.opening_balance('2100-12'),100)      # e 74 anos depois
        self.assertEqual(h.opening_balance('2020-03'),0)        # antes de qualquer dado
        self.assertEqual(h.year_model(2045)['opening'][0],100); self.assertEqual(h.year_model(1999)['opening'][0],0)

    def test_deficit_stops_the_carry_forever(self):
        h=self.make_model(); self.add(h,'2026-01-05','Receita','Receita Marido',100); self.add(h,'2026-03-05','Despesa','Mercado',500,grupo='Essenciais')
        self.assertEqual(h.opening_balance('2026-03'),100); self.assertEqual(h.opening_balance('2026-04'),0); self.assertEqual(h.opening_balance('2040-01'),0)

    def test_get_month_normalizes_and_keeps_last_valid(self):
        h=self.make_model(); h.top_month.set('2045-7'); self.assertEqual(h.get_month(),'2045-07')
        h.top_month.set('lixo'); self.assertEqual(h.get_month(),'2045-07')       # inválido: mantém o último válido
        h.top_month.set('2100-12'); self.assertEqual(h.get_month(),'2100-12')

    def test_month_list_follows_selection_and_data(self):
        h=self.make_model(); h.top_month.set('2060-03'); lst=h.available_months()
        self.assertIn('2060-03',lst); self.assertIn('2066-12',lst); self.assertNotIn('2028-01',lst)
        self.add(h,'2019-02-01','Receita','Receita Marido',1); self.assertIn('2019-02',h.available_months())

if __name__ == '__main__': unittest.main(verbosity=2)
