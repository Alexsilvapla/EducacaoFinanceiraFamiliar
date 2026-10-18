import csv
import math
import shutil
import sqlite3
import tkinter as tk
from copy import copy
from tkinter import ttk, messagebox, filedialog, simpledialog
from pathlib import Path
from datetime import date, datetime

from openpyxl import load_workbook, Workbook
from openpyxl.chart import BarChart as XlBarChart, PieChart as XlPieChart, Reference

APP_NAME = 'Educação Financeira Familiar'
APP_VERSION = '2.1.1'
DB_DIR = Path.home() / 'Documents' / 'EducacaoFinanceiraFamiliar'
DB_DIR.mkdir(parents=True, exist_ok=True)
DB = DB_DIR / 'financas.db'
MONTHS = ['Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
TARGETS_DEFAULT = {'Essenciais':55.0,'Metas':20.0,'Investimentos':10.0,'Gastos livres':10.0,'Educação':5.0}
GROUPS = {
 'Essenciais': list(range(21,32)) + list(range(33,39)) + list(range(40,48)) + list(range(49,53)) + list(range(54,58)) + list(range(59,63)) + list(range(64,68)) + list(range(69,72)) + list(range(73,77)) + list(range(78,81)),
 'Metas': list(range(85,92)), 'Gastos livres': list(range(95,101)),
 'Educação': list(range(104,108)), 'Investimentos': list(range(111,117))}
LESSONS = [
 ('1. Orçamento familiar','Registre receitas e despesas. Compare o planejado com o realizado e revise o orçamento todo mês.'),
 ('2. Necessidades e desejos','Priorize moradia, alimentação, saúde e contas essenciais. Antes de uma compra não essencial, espere e compare preços.'),
 ('3. Dívidas','Liste saldo, juros e vencimento de cada dívida. Priorize juros altos e negocie parcelas que caibam no orçamento.'),
 ('4. Reserva de emergência','Construa gradualmente uma reserva para imprevistos. O tamanho depende das despesas essenciais e da estabilidade da renda.'),
 ('5. Metas financeiras','Defina valor, prazo e aporte mensal de cada objetivo. Separe o dinheiro das metas das despesas do cotidiano.'),
 ('6. Investimentos','Entenda risco, liquidez, custos, impostos e prazo antes de investir. Rentabilidade passada não garante resultados futuros.'),
 ('7. Revisão mensal','Confira o saldo, identifique gastos que cresceram e ajuste o plano sem comprometer necessidades básicas.')]

STRUCTURE = [
 ('Essenciais','TRANSPORTE',['Prestação Carro', 'Seguro Carro', 'Estacionamento', 'Gasolina', 'IPVA', 'Emplacamento', 'Lavação', 'Multas', 'Manutenção Carro', 'Manutenção Moto', 'Outros (transporte)']),
 ('Essenciais','SAÚDE',['Médico', 'Exames', 'Dentista', 'Farmácia', 'Plano de Saúde', 'Outros (saúde)']),
 ('Essenciais','HABITAÇÃO',['Luz', 'Água', 'Crédito celular', 'Gás', 'Canais de TV', 'Internet', 'Manutenção Casa', 'Outros (habitação)']),
 ('Essenciais','ALIMENTAÇÃO',['Mercado', 'Fruteira', 'Padaria', 'Outros (alimentação)']),
 ('Essenciais','CUIDADOS PESSOAIS',['Cabeleireiro', 'Manicure', 'Corte de Cabelo', 'Outros (cuidados pessoais)']),
 ('Essenciais','EDUCAÇÃO BÁSICA',['Material Escolar', 'Uniforme', 'Creche', 'Outros (educação básica)']),
 ('Essenciais','VESTUÁRIO',['Roupas', 'Calçados', 'Acessórios', 'Outros (vestuário)']),
 ('Essenciais','PET',['Cachorro', 'Veterinário', 'Outros (pet)']),
 ('Essenciais','DÍVIDAS',['Lojas', 'Empréstimos', 'Cartão de crédito', 'Outros (dívidas)']),
 ('Essenciais','IGREJA',['Associação Evangelizar', 'Dízimo', 'Ofertas']),
 ('Metas','',['Cama', 'Quarto', 'Muro', 'Garagem', 'Piscina', 'Ar Condicionado', 'Reforma / outros projetos']),
 ('Gastos livres','',['Viagens', 'Cinema', 'Lanches', 'Restaurantes', 'Presentes', 'Outros']),
 ('Educação','',['Pós-Graduação', 'Cursos e certificações', 'Livros', 'Outros']),
 ('Investimentos','',['Ações', 'Tesouro Direto', 'Renda Fixa', 'Fundos de Investimentos', 'Reserva de Emergência', 'Outros']),
]
INCOME_CATS = ['Receita Marido','Receita Esposa','Outras receitas (extras, 13º, férias…)']
OTHER_ROW = 'Outros lançamentos (sem categoria da planilha)'
BLOCKS = ('Essenciais','Metas','Gastos livres','Educação','Investimentos')
BLOCK_LABEL = {'Essenciais':'Essencial','Metas':'Metas','Gastos livres':'Fazer o que quiser','Educação':'Educação','Investimentos':'Investimentos'}
BLOCK_LONG = {'Essenciais':'Essencial','Metas':'Metinhas, Metas e Metonas','Gastos livres':'Fazer o que quiser','Educação':'Educação','Investimentos':'Investimentos'}
LONG_TO_KEY = {v:k for k,v in BLOCK_LONG.items()}
BLOCK_HELP = {
 'Essenciais':'Moradia, alimentação, saúde, transporte, dívidas e demais contas que não dá para deixar de pagar.',
 'Metas':'Metinhas, metas e metonas: reformas, móveis e projetos da casa, com valor e prazo definidos.',
 'Investimentos':'Reserva de emergência e aportes para objetivos de médio e longo prazo.',
 'Gastos livres':'Lazer, viagens, restaurantes e presentes: o que a família quiser, sem culpa e sem estourar o limite.',
 'Educação':'Cursos, livros, certificações e pós-graduação.'}
CATS_BY_GROUP = {'Receitas':set(INCOME_CATS)}
for _g,_sub,_cats in STRUCTURE: CATS_BY_GROUP.setdefault(_g,set()).update(_cats)
_count = {}
for _g,_sub,_cats in STRUCTURE:
    for _c in _cats: _count.setdefault(_c,set()).add(_g)
CAT_GROUP = {c:next(iter(g)) for c,g in _count.items() if len(g)==1}

# ---------------------------------------------------------------- aparência
FONT = 'Segoe UI'
BG = '#f1f5f9'; WHITE = '#ffffff'; BORDER = '#e2e8f0'; TEXT = '#0f172a'; TEXT2 = '#334155'; MUTED = '#64748b'
SIDEBAR = '#0f172a'; SIDEBAR_HOVER = '#1e293b'; SIDEBAR_ACTIVE = '#1d4ed8'
PRIMARY = '#2563eb'; PRIMARY_DARK = '#1d4ed8'
GREEN = '#10b981'; RED = '#ef4444'; AMBER = '#f59e0b'; BLUE = '#2563eb'; PURPLE = '#8b5cf6'; CYAN = '#06b6d4'
PALETTE = ['#2563eb','#10b981','#f59e0b','#ef4444','#8b5cf6','#06b6d4','#ec4899','#84cc16','#f97316','#64748b']
GROUP_COLORS = {'Essenciais':BLUE,'Metas':PURPLE,'Investimentos':GREEN,'Gastos livres':AMBER,'Educação':CYAN}
LABEL_COLORS = {BLOCK_LABEL[k]:c for k,c in GROUP_COLORS.items()}
F_H1 = (FONT,22,'bold'); F_H3 = (FONT,12,'bold'); F_BODY = (FONT,10); F_SMALL = (FONT,9); F_LABEL = (FONT,9,'bold')


def connect():
    con = sqlite3.connect(DB)
    con.execute('''CREATE TABLE IF NOT EXISTS lancamentos (id INTEGER PRIMARY KEY, data TEXT NOT NULL, tipo TEXT NOT NULL, categoria TEXT NOT NULL, descricao TEXT DEFAULT '', valor REAL NOT NULL CHECK(valor>=0), origem TEXT DEFAULT 'manual', grupo TEXT DEFAULT '')''')
    cols = [r[1] for r in con.execute('PRAGMA table_info(lancamentos)')]
    if 'origem' not in cols:
        con.execute("ALTER TABLE lancamentos ADD COLUMN origem TEXT DEFAULT 'manual'")
    if 'grupo' not in cols:
        con.execute("ALTER TABLE lancamentos ADD COLUMN grupo TEXT DEFAULT ''")
    con.execute('''CREATE TABLE IF NOT EXISTS metas (id INTEGER PRIMARY KEY, nome TEXT NOT NULL, valor_alvo REAL NOT NULL, valor_atual REAL NOT NULL DEFAULT 0, prazo TEXT)''')
    con.execute('''CREATE TABLE IF NOT EXISTS preferencias (chave TEXT PRIMARY KEY, valor TEXT NOT NULL)''')
    for k,v in TARGETS_DEFAULT.items(): con.execute('INSERT OR IGNORE INTO preferencias(chave,valor) VALUES(?,?)',(f'alvo_{k}',str(v)))
    con.execute('''CREATE TABLE IF NOT EXISTS saldos_iniciais (mes TEXT PRIMARY KEY, valor REAL NOT NULL DEFAULT 0)''')
    con.commit()
    return con


def extract_workbook_data(path):
    """Read only manually entered monthly values from the source workbook."""
    wb=load_workbook(path,data_only=True,read_only=True)
    try:
        if 'Orçamento Mensal' not in wb.sheetnames:
            raise ValueError('Não encontrei a aba “Orçamento Mensal”.')
        ws=wb['Orçamento Mensal']; year=int(ws['D2'].value or date.today().year)
        entries=[]
        for row in (5,6,7):
            cat=str(ws.cell(row,3).value or 'Receita')
            for col in range(4,16):
                value=num(ws.cell(row,col).value)
                if value>0: entries.append((f'{year}-{col-3:02d}-01','Receita',cat,value,''))
        row_to_type={r:'Despesa' for r in GROUPS['Essenciais']+GROUPS['Metas']+GROUPS['Gastos livres']+GROUPS['Educação']}
        row_to_type.update({r:'Investimento' for r in GROUPS['Investimentos']})
        for row,typ in row_to_type.items():
            cat=str(ws.cell(row,3).value or f'Categoria da linha {row}')
            group=next((g for g,rows in GROUPS.items() if row in rows),'')
            for col in range(4,16):
                value=num(ws.cell(row,col).value)
                if value>0: entries.append((f'{year}-{col-3:02d}-01',typ,cat,value,group))
        targets={}
        for row,key in {13:'Essenciais',14:'Metas',15:'Investimentos',16:'Gastos livres',17:'Educação'}.items():
            raw=ws.cell(row,2).value
            if isinstance(raw,(int,float)): targets[key]=float(raw)*100 if raw<=1 else float(raw)
        opening=num(ws['D9'].value)
        return year,entries,opening,targets
    finally:
        wb.close()

def money(v):
    return ('R$ {:,.2f}'.format(float(v or 0))).replace(',','X').replace('.',',').replace('X','.')


def parse_amount(value):
    if value is None or value == '': return 0.0
    if isinstance(value, (int,float)): return float(value)
    s = str(value).strip().replace('R$','').replace(' ','')
    if ',' in s: s=s.replace('.','').replace(',','.')
    return float(s)


def num(value):
    try: return float(value or 0)
    except (ValueError, TypeError): return 0.0


# ---------------------------------------------------------------- utilitários dos gráficos
def safe_amount(value):
    try: return parse_amount(value)
    except (ValueError, TypeError): return 0.0


def short_money(v):
    """Formato curto para eixos: 1,5 mil / 2 mi."""
    v=float(v or 0); a=abs(v)
    if a>=1e6: s=f'{v/1e6:.1f} mi'
    elif a>=1e3: s=f'{v/1e3:.1f} mil'
    else: s=f'{v:.0f}'
    return s.replace('.',',').replace(',0 ',' ')


def nice_max(v):
    """Topo do eixo escolhido para que as 4 divisões sejam números redondos (ex.: 6 mil, 12 mil, 18 mil, 24 mil)."""
    if v<=0: return 1.0
    raw=v/4; mag=10**math.floor(math.log10(raw))
    for m in (1,1.2,1.5,2,2.5,3,4,5,6,8,10):
        if m*mag>=raw-1e-12: return round(4*m*mag,10)
    return 40*mag


def top_n(items,n,other='Outras'):
    """items: [(rótulo, valor)]. Mantém os n-1 maiores e agrupa o resto em 'Outras'."""
    items=sorted([(l,v) for l,v in items if v>0],key=lambda x:x[1],reverse=True)
    if len(items)<=n: return items
    rest=sum(v for _,v in items[n-1:])
    return items[:n-1]+[(other,rest)]


def colorize(items,fixed=None):
    fixed=fixed or {}
    return [(l,v,fixed.get(l,PALETTE[i%len(PALETTE)])) for i,(l,v) in enumerate(items)]


def month_label(month):
    try: y,m=map(int,month.split('-')); return f'{MONTHS[m-1]} de {y}'
    except (ValueError, IndexError): return month


def months_until(prazo):
    try:
        py,pm=map(int,prazo.split('-')); t=date.today(); return (py-t.year)*12+pm-t.month
    except (ValueError, AttributeError): return None


def blk(g):
    return BLOCK_LABEL.get(g,g)


def eff_group(tipo,grupo,cat):
    """Bloco efetivo de um lançamento, como a planilha o classificaria."""
    if tipo=='Receita': return 'Receitas'
    if tipo=='Investimento': return 'Investimentos'
    return grupo or CAT_GROUP.get(cat) or 'Essenciais'


def br(v):
    """Número no padrão brasileiro, sem R$ (para a planilha); zero vira traço."""
    return '–' if abs(v)<0.005 else '{:,.2f}'.format(v).replace(',','X').replace('.',',').replace('X','.')


def pctf(v):
    return f'{v:.0%}' if abs(v)>=0.0005 else '–'


def split_income(renda,targets):
    """Divide a renda pelos percentuais-alvo: [(bloco, percentual, valor)]."""
    return [(g,p,renda*p/100) for g,p in targets.items()]


def reserve_target(despesas_mensais,meses):
    return max(0.0,despesas_mensais)*max(0.0,meses)


def months_to_goal(valor,aporte):
    if valor<=0: return 0
    if aporte<=0: return None
    return math.ceil(valor/aporte-1e-9)


def duration_text(n):
    a,m=divmod(int(n),12); parts=[]
    if a: parts.append(f'{a} {"ano" if a==1 else "anos"}')
    if m or not a: parts.append(f'{m} {"mês" if m==1 else "meses"}')
    return ' e '.join(parts)


def compound_projection(aporte,taxa_anual,anos,inicial=0.0):
    """Simulação ilustrativa: aporte mensal com taxa constante. Devolve [(ano, investido, saldo)]."""
    i=(1+max(0.0,taxa_anual)/100)**(1/12)-1; bal=float(inicial); invested=float(inicial); out=[]
    for ano in range(1,int(anos)+1):
        for _ in range(12):
            bal=bal*(1+i)+aporte; invested+=aporte
        out.append((ano,invested,bal))
    return out


def round_rect(c,x0,y0,x1,y1,r=6,**kw):
    r=max(0,min(r,(x1-x0)/2,(y1-y0)/2))
    pts=[x0+r,y0,x1-r,y0,x1,y0,x1,y0+r,x1,y1-r,x1,y1,x1-r,y1,x0+r,y1,x0,y1,x0,y1-r,x0,y0+r,x0,y0]
    return c.create_polygon(pts,smooth=True,**kw)


def card(parent,title=None,subtitle=None):
    f=tk.Frame(parent,bg=WHITE,highlightbackground=BORDER,highlightthickness=1)
    if title: tk.Label(f,text=title,font=F_H3,bg=WHITE,fg=TEXT).pack(anchor='w',padx=16,pady=(14,0 if subtitle else 8))
    if subtitle: tk.Label(f,text=subtitle,font=F_SMALL,bg=WHITE,fg=MUTED).pack(anchor='w',padx=16,pady=(0,8))
    return f


def button(parent,text,command,kind='primary'):
    colors={'primary':(PRIMARY,'white',PRIMARY_DARK),'ghost':('#e2e8f0',TEXT2,'#cbd5e1'),'danger':('#fee2e2','#b91c1c','#fecaca')}
    bg,fg,hover=colors[kind]
    b=tk.Button(parent,text=text,command=command,bg=bg,fg=fg,activebackground=hover,activeforeground=fg,font=(FONT,10,'bold'),bd=0,relief='flat',padx=16,pady=8,cursor='hand2')
    b.bind('<Enter>',lambda e:b.configure(bg=hover)); b.bind('<Leave>',lambda e:b.configure(bg=bg))
    return b


class ChartBase(tk.Canvas):
    """Canvas que se redesenha ao mudar de tamanho e mostra dica (tooltip) ao passar o mouse."""
    def __init__(self,parent,height=240):
        super().__init__(parent,bg=WHITE,height=height,highlightthickness=0)
        self.bind('<Configure>',lambda e:self.redraw())

    def redraw(self): pass

    def hover(self,tag,text):
        self.tag_bind(tag,'<Motion>',lambda e,t=text:self.tip(e,t))
        self.tag_bind(tag,'<Leave>',lambda e:self.delete('tip'))

    def tip(self,event,text):
        self.delete('tip')
        t=self.create_text(event.x+16,event.y+14,text=text,anchor='nw',font=F_SMALL,fill='white',tags='tip')
        x0,y0,x1,y1=self.bbox(t); w,h=self.winfo_width(),self.winfo_height(); dx=dy=0
        if x1+10>w: dx=-(x1+10-w)-24
        if y1+10>h: dy=-(y1-y0)-30
        self.move('tip',dx,dy)
        x0,y0,x1,y1=self.bbox(t)
        r=round_rect(self,x0-8,y0-6,x1+8,y1+6,6,fill=TEXT,outline='',tags='tip')
        self.tag_lower(r,t)


class DonutChart(ChartBase):
    """Gráfico de pizza (donut=False) ou rosca (donut=True) com legenda e dica ao passar o mouse."""
    def __init__(self,parent,donut=True,fmt=None,center_title='Total',height=230):
        super().__init__(parent,height=height)
        self.donut=donut; self.fmt=fmt or money; self.center_title=center_title; self.slices=[]

    def set_data(self,items):
        self.slices=[(l,float(v),c) for l,v,c in items if v>0]; self.redraw()

    def redraw(self):
        self.delete('all'); w,h=self.winfo_width(),self.winfo_height()
        if w<60 or h<60: return
        wide=w>=340
        size=max(80,min(h-24,int(w*0.48)) if wide else min(w-28,h-24)); x0=14 if wide else (w-size)/2; y0=(h-size)/2
        x1,y1=x0+size,y0+size; cx,cy=x0+size/2,y0+size/2; thick=max(22,size*0.20)
        total=sum(v for _,v,_ in self.slices)
        if total<=0:
            self.create_oval(x0+thick/2,y0+thick/2,x1-thick/2,y1-thick/2,outline=BORDER,width=thick) if self.donut else self.create_oval(x0,y0,x1,y1,fill=BG,outline=BORDER)
            self.create_text(cx,cy,text='Sem dados\nneste mês',font=F_SMALL,fill=MUTED,justify='center'); return
        start=90
        for i,(label,val,col) in enumerate(self.slices):
            ext=-val/total*360
            if len(self.slices)==1: ext=-359.99
            tag=f's{i}'
            if self.donut:
                p=thick/2; self.create_arc(x0+p,y0+p,x1-p,y1-p,start=start,extent=ext,style='arc',width=thick,outline=col,tags=tag)
            else:
                self.create_arc(x0,y0,x1,y1,start=start,extent=ext,style='pieslice',fill=col,outline=WHITE,width=2,tags=tag)
            self.hover(tag,f'{label}\n{self.fmt(val)}  ·  {val/total:.1%}'); start+=ext
        if self.donut:
            self.create_text(cx,cy-9,text=self.center_title,font=F_SMALL,fill=MUTED)
            txt=self.fmt(total); self.create_text(cx,cy+10,text=txt,font=(FONT,10 if len(txt)>10 else 13,'bold'),fill=TEXT)
        if wide:
            lx=x1+28; avail=w-lx-12; n=len(self.slices); rowh=min(26,(h-16)/max(n,1)); ly=(h-rowh*n)/2; maxc=max(6,int((avail-48)/6.6))
            for i,(label,val,col) in enumerate(self.slices):
                y=ly+i*rowh+rowh/2; name=label if len(label)<=maxc else label[:maxc-1]+'…'
                round_rect(self,lx,y-5,lx+10,y+5,3,fill=col,outline='')
                self.create_text(lx+18,y,text=name,anchor='w',font=F_BODY,fill=TEXT2)
                self.create_text(w-12,y,text=f'{val/total:.0%}',anchor='e',font=(FONT,10,'bold'),fill=TEXT)


class BarChart(ChartBase):
    """Barras agrupadas com grade, eixo em R$ e dica com os valores de cada mês."""
    def __init__(self,parent,height=250):
        super().__init__(parent,height=height); self.labels=[]; self.series=[]

    def set_data(self,labels,series):
        self.labels=labels; self.series=series; self.redraw()

    def redraw(self):
        self.delete('all'); w,h=self.winfo_width(),self.winfo_height()
        if w<140 or h<110: return
        left,right,top,bottom=58,14,34,h-28; ph=bottom-top
        mx=nice_max(max([v for _,_,vals in self.series for v in vals]+[0]))
        for i in range(5):
            y=bottom-ph*i/4
            self.create_line(left,y,w-right,y,fill=BORDER if i else '#cbd5e1')
            self.create_text(left-8,y,text=short_money(mx*i/4),anchor='e',font=F_SMALL,fill=MUTED)
        n=len(self.labels); k=max(len(self.series),1); gw=(w-left-right)/max(n,1); bw=max(6,min(26,gw*0.72/k)); lx=left
        for name,col,_ in self.series:
            round_rect(self,lx,12,lx+10,22,3,fill=col,outline=''); self.create_text(lx+16,17,text=name,anchor='w',font=F_SMALL,fill=TEXT2); lx+=len(name)*6.4+42
        for j,lab in enumerate(self.labels):
            gx=left+j*gw+gw/2; start=gx-bw*k/2; tag=f'g{j}'
            self.create_rectangle(left+j*gw,top,left+(j+1)*gw,bottom,fill=WHITE,outline='',tags=tag)
            for s,(name,col,vals) in enumerate(self.series):
                bh=vals[j]/mx*ph
                if bh>0.5: round_rect(self,start+s*bw+1,bottom-bh,start+(s+1)*bw-1,bottom,3,fill=col,outline='',tags=tag)
            self.create_text(gx,bottom+14,text=lab,font=F_SMALL,fill=MUTED,tags=tag)
            self.hover(tag,lab+'\n'+'\n'.join(f'{name}: {money(vals[j])}' for name,_,vals in self.series))
        if not any(v for _,_,vals in self.series for v in vals):
            self.create_text((left+w)/2,(top+bottom)/2,text='Sem lançamentos neste período',font=F_BODY,fill=MUTED)


class HBarList(ChartBase):
    """Lista de barras de progresso horizontais (orçamento x realizado, metas)."""
    ROW_H=48
    def __init__(self,parent,fmt=None):
        super().__init__(parent,height=60); self.rows=[]; self.fmt=fmt or money

    def set_rows(self,rows):
        """rows: [(rótulo, usado, total, cor ou None)]"""
        self.rows=rows; self.configure(height=max(60,len(rows)*self.ROW_H+10)); self.redraw()

    def redraw(self):
        self.delete('all'); w=self.winfo_width()
        if w<100: return
        if not self.rows:
            self.create_text(w/2,30,text='Nada para exibir ainda',font=F_BODY,fill=MUTED); return
        pad=4
        for i,(label,used,total,color) in enumerate(self.rows):
            y=i*self.ROW_H+6; ratio=(used/total) if total>0 else (1.0 if used>0 else 0.0)
            col=color or (RED if ratio>1 else AMBER if ratio>=.85 else GREEN)
            self.create_text(pad,y+8,text=label,anchor='w',font=(FONT,10,'bold'),fill=TEXT)
            txt=f'{self.fmt(used)} de {self.fmt(total)}  ·  {ratio:.0%}' if total>0 else f'{self.fmt(used)}  ·  sem orçamento'
            self.create_text(w-pad,y+8,text=txt,anchor='e',font=F_SMALL,fill=RED if ratio>1 and not color else MUTED)
            round_rect(self,pad,y+22,w-pad,y+32,5,fill=BORDER,outline='')
            fw=(w-2*pad)*min(ratio,1)
            if fw>0.5: round_rect(self,pad,y+22,pad+max(fw,10),y+32,5,fill=col,outline='')


class App(tk.Tk):
    NAV=[('Painel','▦'),('Painel do ano','◧'),('Planilha','▥'),('Lançamentos','＋'),('Orçamento','◫'),('Metas','◎'),('Investimentos','↗'),('Relatórios','▤'),('Aprender','✦')]

    def __init__(self):
        super().__init__()
        self.title(f'{APP_NAME} · {APP_VERSION}')
        self.geometry('1320x820'); self.minsize(1100,700); self.configure(bg=BG)
        self.con=connect(); self.current_page='Painel'; self.nav_buttons={}
        self._setup_style(); self._build_shell(); self.show_page('Painel')
        self.bind_all('<MouseWheel>',self._on_wheel)
        self.protocol('WM_DELETE_WINDOW', self.close)

    # ------------------------------------------------------------ estilo e estrutura
    def _setup_style(self):
        s=ttk.Style(self); s.theme_use('clam'); self.style=s
        s.configure('.',font=(FONT,10))
        s.configure('Treeview',font=(FONT,10),rowheight=31,background=WHITE,fieldbackground=WHITE,foreground=TEXT2,borderwidth=0)
        s.configure('Treeview.Heading',font=(FONT,9,'bold'),background='#f8fafc',foreground=MUTED,relief='flat',padding=(8,8))
        s.map('Treeview',background=[('selected','#dbeafe')],foreground=[('selected',TEXT)])
        s.map('Treeview.Heading',background=[('active','#f1f5f9')])
        s.configure('TEntry',padding=6,fieldbackground=WHITE); s.configure('TCombobox',padding=6,fieldbackground=WHITE)
        s.configure('Vertical.TScrollbar',background='#cbd5e1',troughcolor=BG,borderwidth=0,arrowsize=12)

    def _build_shell(self):
        shell=tk.Frame(self,bg=BG); shell.pack(fill='both',expand=True)
        sb=self.sidebar=tk.Frame(shell,bg=SIDEBAR,width=240); sb.pack(side='left',fill='y'); sb.pack_propagate(False)
        logo=tk.Frame(sb,bg=SIDEBAR); logo.pack(fill='x',padx=20,pady=(20,16))
        tk.Label(logo,text='EF',font=(FONT,16,'bold'),bg=PRIMARY,fg='white',width=3,pady=6).pack(side='left')
        tk.Label(logo,text='Educação\nFinanceira',font=(FONT,11,'bold'),bg=SIDEBAR,fg='white',justify='left').pack(side='left',padx=12)
        tk.Label(sb,text='MENU',font=(FONT,8,'bold'),bg=SIDEBAR,fg='#64748b').pack(anchor='w',padx=22,pady=(0,6))
        for label,icon in self.NAV:
            b=tk.Button(sb,text=f'  {icon}   {label}',anchor='w',font=(FONT,11),bd=0,relief='flat',padx=16,pady=6,bg=SIDEBAR,fg='#cbd5e1',activebackground=SIDEBAR_HOVER,activeforeground='white',cursor='hand2',command=lambda x=label:self.show_page(x))
            b.pack(fill='x',padx=12,pady=2); self.nav_buttons[label]=b
            b.bind('<Enter>',lambda e,x=label:self._nav_hover(x,True)); b.bind('<Leave>',lambda e,x=label:self._nav_hover(x,False))
        tk.Frame(sb,bg='#1e293b',height=1).pack(fill='x',padx=20,pady=12)
        for text,cmd,bg,hov in [('  ⇧   Importar planilha Excel',self.import_workbook,PRIMARY,PRIMARY_DARK),('  ⤓   Backup / restauração',self.backup_menu,SIDEBAR_HOVER,'#334155'),('  ⟲   Restaurar do zero',self.reset_all,'#7f1d1d','#991b1b')]:
            b=tk.Button(sb,text=text,anchor='w',font=(FONT,10,'bold'),bd=0,padx=16,pady=7,bg=bg,fg='white',activebackground=hov,activeforeground='white',cursor='hand2',command=cmd)
            b.pack(fill='x',padx=12,pady=3); b.bind('<Enter>',lambda e,b=b,h=hov:b.configure(bg=h)); b.bind('<Leave>',lambda e,b=b,c=bg:b.configure(bg=c))
        tk.Label(sb,text=f'Versão {APP_VERSION}\nDados salvos localmente\nno seu computador',font=F_SMALL,bg=SIDEBAR,fg='#64748b',justify='left').pack(side='bottom',anchor='w',padx=22,pady=12)

        self.main=tk.Frame(shell,bg=BG); self.main.pack(side='left',fill='both',expand=True)
        top=tk.Frame(self.main,bg=WHITE,height=64,highlightbackground=BORDER,highlightthickness=1); top.pack(fill='x'); top.pack_propagate(False)
        self.top_title=tk.Label(top,text='Painel',font=(FONT,16,'bold'),bg=WHITE,fg=TEXT); self.top_title.pack(side='left',padx=28)
        nav=tk.Frame(top,bg=WHITE); nav.pack(side='right',padx=24)
        self.top_month=tk.StringVar(value=date.today().strftime('%Y-%m'))
        self.month_lbl=tk.Label(nav,text='',font=(FONT,10,'bold'),bg=WHITE,fg=TEXT2,width=20,anchor='e'); self.month_lbl.pack(side='left',padx=(0,10))
        for sym,d in (('«',-12),('◀',-1)): b_=button(nav,sym,lambda d=d:self.shift_month(d),'ghost'); b_.configure(padx=10,pady=4); b_.pack(side='left',padx=2)
        self.month_combo=ttk.Combobox(nav,textvariable=self.top_month,values=self.available_months(),width=9,state='normal'); self.month_combo.pack(side='left',padx=2)
        for sym,d in (('▶',1),('»',12)): b_=button(nav,sym,lambda d=d:self.shift_month(d),'ghost'); b_.configure(padx=10,pady=4); b_.pack(side='left',padx=2)
        self.month_combo.bind('<<ComboboxSelected>>',lambda e:self.refresh_page()); self.month_combo.bind('<Return>',lambda e:self.refresh_page()); self.month_combo.bind('<FocusOut>',lambda e:self.refresh_page())

        # área de conteúdo com rolagem
        wrap=tk.Frame(self.main,bg=BG); wrap.pack(fill='both',expand=True)
        self.scroll_canvas=tk.Canvas(wrap,bg=BG,highlightthickness=0)
        vbar=ttk.Scrollbar(wrap,orient='vertical',command=self.scroll_canvas.yview); self.scroll_canvas.configure(yscrollcommand=vbar.set)
        vbar.pack(side='right',fill='y'); self.scroll_canvas.pack(side='left',fill='both',expand=True)
        self.content=tk.Frame(self.scroll_canvas,bg=BG)
        self._win=self.scroll_canvas.create_window((0,0),window=self.content,anchor='nw')
        self.content.bind('<Configure>',lambda e:self.scroll_canvas.configure(scrollregion=self.scroll_canvas.bbox('all')))
        self.scroll_canvas.bind('<Configure>',lambda e:self.scroll_canvas.itemconfigure(self._win,width=e.width-48))
        self.scroll_canvas.move(self._win,24,0)

    def _nav_hover(self,name,inside):
        if name==self.current_page: return
        self.nav_buttons[name].configure(bg=SIDEBAR_HOVER if inside else SIDEBAR)

    def _on_wheel(self,event):
        if isinstance(event.widget,(ttk.Treeview,ttk.Combobox)): return
        if self.scroll_canvas.yview()==(0.0,1.0): return
        self.scroll_canvas.yview_scroll(int(-event.delta/120),'units')

    def available_months(self):
        cur=self.get_month() if hasattr(self,'top_month') else date.today().strftime('%Y-%m'); y=int(cur[:4])
        lo,hi=self.con.execute("SELECT MIN(substr(data,1,4)),MAX(substr(data,1,4)) FROM lancamentos").fetchone()
        first=max(1,min([y-6]+([int(lo)] if lo else []))); last=min(9999,max([y+6]+([int(hi)] if hi else [])))
        return [f'{yy}-{m:02d}' for yy in range(first,last+1) for m in range(1,13)]

    def shift_month(self,delta):
        y,m=map(int,self.get_month().split('-')); idx=y*12+(m-1)+delta
        self.top_month.set(f'{idx//12}-{idx%12+1:02d}'); self.refresh_page()

    def show_page(self,name):
        self.current_page=name; self.top_title.configure(text=name)
        for k,b in self.nav_buttons.items(): b.configure(bg=SIDEBAR_ACTIVE if k==name else SIDEBAR,fg='white' if k==name else '#cbd5e1')
        for w in self.content.winfo_children(): w.destroy()
        self.scroll_canvas.yview_moveto(0)
        {'Painel':self.build_dashboard,'Painel do ano':self.build_year,'Planilha':self.build_grid,'Lançamentos':self.build_mov,'Orçamento':self.build_budget,'Metas':self.build_goals,'Investimentos':self.build_investments,'Relatórios':self.build_reports,'Aprender':self.build_lessons}[name]()
        self.refresh_page()

    # ------------------------------------------------------------ helpers de layout
    def heading(self,title,subtitle=None):
        tk.Label(self.content,text=title,font=F_H1,bg=BG,fg=TEXT).pack(anchor='w',pady=(22,0))
        tk.Label(self.content,text=subtitle or '',font=F_BODY,bg=BG,fg=MUTED).pack(anchor='w',pady=(4,16))

    def cols(self,*weights,pady=(0,14)):
        f=tk.Frame(self.content,bg=BG); f.pack(fill='x',pady=pady); f.rowconfigure(0,weight=1)
        for i,wt in enumerate(weights): f.columnconfigure(i,weight=wt,uniform='c')
        f._n=len(weights); return f

    def put(self,frame,widget,col):
        widget.grid(row=0,column=col,sticky='nsew',padx=(0 if col==0 else 7,0 if col==frame._n-1 else 7))

    def table(self,parent,columns,headings,widths,height=10,right=()):
        outer=tk.Frame(parent,bg=WHITE,highlightbackground=BORDER,highlightthickness=1)
        tree=ttk.Treeview(outer,columns=columns,show='headings',height=height)
        for c,h,w in zip(columns,headings,widths):
            a='e' if c in right else 'w'; tree.heading(c,text=h,anchor=a); tree.column(c,width=w,anchor=a)
        vs=ttk.Scrollbar(outer,orient='vertical',command=tree.yview); tree.configure(yscrollcommand=vs.set)
        tree.tag_configure('odd',background='#f8fafc'); tree.tag_configure('even',background=WHITE)
        tree.tag_configure('Receita',foreground='#047857'); tree.tag_configure('Despesa',foreground='#b91c1c'); tree.tag_configure('Investimento',foreground='#1d4ed8')
        tree.pack(side='left',fill='both',expand=True); vs.pack(side='right',fill='y')
        return tree

    def fill(self,tree,rows):
        """rows: [(iid ou None, valores, tag ou None)]"""
        for i in tree.get_children(): tree.delete(i)
        for n,(iid,values,tag) in enumerate(rows):
            kw={'values':values,'tags':((tag,) if tag else ())+(('odd' if n%2 else 'even'),)}
            if iid is not None: kw['iid']=str(iid)
            tree.insert('','end',**kw)

    def labeled_entry(self,parent,text,var,width,col,row=0,widget='entry',values=None):
        tk.Label(parent,text=text,font=F_LABEL,bg=WHITE,fg=MUTED).grid(row=row,column=col,sticky='w',padx=10,pady=(12,2))
        if widget=='combo': w=ttk.Combobox(parent,textvariable=var,values=values,width=width,state='readonly')
        elif widget=='combo_free': w=ttk.Combobox(parent,textvariable=var,values=values,width=width)
        else: w=ttk.Entry(parent,textvariable=var,width=width)
        w.grid(row=row+1,column=col,sticky='ew',padx=10,pady=(0,12)); return w

    # ------------------------------------------------------------ dados
    def get_month(self):
        value=self.top_month.get().strip()
        try: dt=datetime.strptime(value,'%Y-%m'); self._last_month=f'{dt.year:04d}-{dt.month:02d}'; return self._last_month
        except ValueError: return getattr(self,'_last_month',date.today().strftime('%Y-%m'))

    def month_rows(self,month=None):
        month=month or self.get_month()
        return self.con.execute("SELECT id,data,tipo,categoria,descricao,valor,origem,grupo FROM lancamentos WHERE substr(data,1,7)=? ORDER BY data DESC,id DESC",(month,)).fetchall()

    def sums(self,month=None):
        result={'Receita':0.0,'Despesa':0.0,'Investimento':0.0}
        for r in self.month_rows(month): result[r[2]]=result.get(r[2],0)+r[5]
        return result

    def by_category(self,month,tipo):
        out={}
        for r in self.month_rows(month):
            if r[2]==tipo: out[r[3]]=out.get(r[3],0)+r[5]
        return out

    def last_months(self,month,count=6):
        y,m=map(int,month.split('-')); out=[]
        for i in range(count-1,-1,-1):
            idx=y*12+(m-1)-i; out.append(f'{idx//12}-{idx%12+1:02d}')
        return out

    def targets(self):
        return {k:parse_amount(self.pref(f'alvo_{k}',v)) for k,v in TARGETS_DEFAULT.items()}

    def opening_balance(self,month=None):
        """Sobra do mês anterior, como na planilha: o saldo positivo de um mês vira a sobra do seguinte.
        Calculada a partir do primeiro mês com dados (ou sobra informada), sem limite de anos."""
        month=month or self.get_month()
        try: target=int(month[:4])*12+int(month[5:7])-1
        except ValueError: return 0.0
        anchors=dict(self.con.execute('SELECT mes,valor FROM saldos_iniciais').fetchall()); flow={}
        for mes,tipo,v in self.con.execute('SELECT substr(data,1,7),tipo,SUM(valor) FROM lancamentos GROUP BY substr(data,1,7),tipo'):
            flow[mes]=flow.get(mes,0.0)+(v if tipo=='Receita' else -v)
        keys=[k for k in list(flow)+list(anchors) if len(k)==7]
        if not keys: return 0.0
        idx=lambda k:int(k[:4])*12+int(k[5:7])-1; start=min(idx(k) for k in keys)
        if target<start: return 0.0
        carry=0.0
        for i in range(start,target+1):
            mes=f'{i//12:04d}-{i%12+1:02d}'; opening=anchors.get(mes,carry)
            if i==target: return opening
            carry=max(0.0,opening+flow.get(mes,0.0))
        return 0.0

    def pref(self,key,default):
        r=self.con.execute('SELECT valor FROM preferencias WHERE chave=?',(key,)).fetchone(); return r[0] if r else str(default)

    def group_spend(self,month):
        result={k:0.0 for k in TARGETS_DEFAULT}
        for tipo,cat,val,stored_group in self.con.execute("SELECT tipo,categoria,valor,grupo FROM lancamentos WHERE substr(data,1,7)=?",(month,)):
            if tipo=='Receita': continue
            group=stored_group or self.category_group(cat)
            if tipo=='Investimento': group='Investimentos'
            if group: result[group]+=val
        return result

    def category_group(self,cat):
        if cat in CAT_GROUP: return CAT_GROUP[cat]
        mapping={'Casa e reformas':'Metas','Objetivos familiares':'Metas','Outros projetos':'Metas','Viagens':'Gastos livres','Cinema':'Gastos livres','Cinema e lazer':'Gastos livres','Lanches':'Gastos livres','Lanches e restaurantes':'Gastos livres','Restaurantes':'Gastos livres','Presentes':'Gastos livres','Cursos e certificações':'Educação','Livros':'Educação','Pós-Graduação':'Educação','Graduação/Pós-graduação':'Educação','Ações':'Investimentos','Tesouro Direto':'Investimentos','Renda Fixa':'Investimentos','Renda fixa':'Investimentos','Fundos de Investimentos':'Investimentos','Reserva de Emergência':'Investimentos','Reserva de emergência':'Investimentos'}
        if cat in mapping: return mapping[cat]
        if cat in ('Cama','Quarto','Muro','Garagem','Piscina','Ar Condicionado','Reforma / outros projetos'): return 'Metas'
        if cat in ('Viagens','Cinema','Lanches','Restaurantes','Outros') and cat=='Outros': return 'Gastos livres'
        if cat in ('Receita Marido','Receita Esposa','Outras receitas','Outras receitas (extras, 13º, férias…)'): return None
        if cat in ('Ações','Tesouro Direto','Renda Fixa','Fundos de Investimentos','Reserva de Emergência'): return 'Investimentos'
        # The type is considered by group_spend; general categories default to essentials.
        return 'Essenciais'

    def all_categories(self):
        return ['Receita Marido','Receita Esposa','Outras receitas'] + ['Transporte','Saúde','Habitação','Alimentação','Cuidados pessoais','Educação básica','Vestuário','Pet','Dívidas','Igreja','Casa e reformas','Objetivos familiares','Outros projetos','Viagens','Cinema e lazer','Lanches e restaurantes','Presentes','Cursos e certificações','Livros','Graduação/Pós-graduação','Ações','Tesouro Direto','Renda fixa','Fundos de investimentos','Reserva de emergência','Outros']

    def year_model(self,year):
        """Recalcula o que a aba «Orçamento Mensal» calcula, mês a mês, para o ano."""
        y=str(year); cell={}
        q="SELECT tipo,grupo,categoria,substr(data,6,2),SUM(valor) FROM lancamentos WHERE substr(data,1,4)=? GROUP BY tipo,grupo,categoria,substr(data,6,2)"
        for tipo,grupo,cat,mes,val in self.con.execute(q,(y,)):
            g=eff_group(tipo,grupo,cat); key=(g,cat if cat in CATS_BY_GROUP.get(g,()) else OTHER_ROW)
            cell.setdefault(key,[0.0]*12)[int(mes)-1]+=val
        def tot(g): return [sum(v[i] for (gg,_),v in cell.items() if gg==g) for i in range(12)]
        rec=tot('Receitas'); inv=tot('Investimentos'); desp_g={g:tot(g) for g in ('Essenciais','Metas','Gastos livres','Educação')}
        despesas=[sum(desp_g[g][i] for g in desp_g) for i in range(12)]; saidas=[despesas[i]+inv[i] for i in range(12)]
        anchors=dict(self.con.execute('SELECT mes,valor FROM saldos_iniciais').fetchall())
        opening=[0.0]*12; saldo=[0.0]*12
        for i in range(12):
            mes=f'{y}-{i+1:02d}'
            if mes in anchors: opening[i]=anchors[mes]
            elif i==0: opening[i]=self.opening_balance(mes)
            else: opening[i]=max(0.0,saldo[i-1])
            saldo[i]=opening[i]+rec[i]-saidas[i]
        targets=self.targets(); orcado={g:[rec[i]*targets.get(g,0)/100 for i in range(12)] for g in targets}
        pat=[]; acc=0.0
        for v in inv: acc+=v; pat.append(acc)
        return dict(cell=cell,rec=rec,inv=inv,desp_g=desp_g,despesas=despesas,saidas=saidas,opening=opening,saldo=saldo,orcado=orcado,patrimonio=pat,targets=targets)

    @staticmethod
    def block_series(m,g):
        return m['inv'] if g=='Investimentos' else m['desp_g'][g]

    def grid_rows(self,m,year):
        """Linhas da aba «Orçamento Mensal»: [(iid, rótulo, 12 células, ano, tag, chave de edição)]."""
        rows=[]; cell=m['cell']; zeros=[0.0]*12; S=sum; n=[0]
        def add(label,vals,ano,tag,edit=None,fmt=br):
            n[0]+=1; rows.append((f'r{n[0]}',label,['']*12 if vals is None else [fmt(v) for v in vals],'' if ano is None else fmt(ano),tag,edit))
        add('1. RECEITAS',None,None,'sec')
        for c in INCOME_CATS: v=cell.get(('Receitas',c),zeros); add('    '+c,v,S(v),'item',('Receitas',c))
        if ('Receitas',OTHER_ROW) in cell: v=cell[('Receitas',OTHER_ROW)]; add('    '+OTHER_ROW,v,S(v),'item')
        add('RECEITA DO MÊS (base do rateio)',m['rec'],S(m['rec']),'tot')
        add('Sobra do mês anterior  (duplo clique em janeiro para informar)',m['opening'],None,'item','opening')
        add('TOTAL DISPONÍVEL NO MÊS',[m['rec'][i]+m['opening'][i] for i in range(12)],S(m['rec'])+m['opening'][0],'tot')
        add('2. DIVISÃO DA RECEITA — QUANTO PODE GASTAR EM CADA BLOCO (ORÇADO)',None,None,'sec')
        for g,p in m['targets'].items(): add(f'    {blk(g)} ({p:g}%)',m['orcado'][g],S(m['orcado'][g]),'item')
        titles={'Essenciais':'3. DESPESAS ESSENCIAIS','Metas':'4. METINHAS, METAS E METONAS','Gastos livres':'5. FAZER O QUE QUISER','Educação':'6. EDUCAÇÃO','Investimentos':'7. INVESTIMENTOS'}
        totals={'Essenciais':'TOTAL ESSENCIAL','Metas':'TOTAL METAS','Gastos livres':'TOTAL FAZER O QUE QUISER','Educação':'TOTAL EDUCAÇÃO','Investimentos':'TOTAL INVESTIDO'}
        for g in BLOCKS:
            add(f'{titles[g]}  (bloco de {m["targets"].get(g,0):g}%)',None,None,'sec')
            for gg,sub,cats in STRUCTURE:
                if gg!=g: continue
                if sub: add('  '+sub.capitalize(),None,None,'sub')
                for c in cats: v=cell.get((g,c),zeros); add('    '+c,v,S(v),'item',(g,c))
                if sub:
                    sv=[S(cell.get((g,c),zeros)[i] for c in cats) for i in range(12)]; add('  Total '+sub.capitalize(),sv,S(sv),'subtot')
            if (g,OTHER_ROW) in cell: v=cell[(g,OTHER_ROW)]; add('    '+OTHER_ROW,v,S(v),'item')
            ser=self.block_series(m,g); add(totals[g],ser,S(ser),'tot')
            if g=='Investimentos': add('Patrimônio acumulado (soma dos aportes)',m['patrimonio'],m['patrimonio'][-1],'item')
        rec_y=S(m['rec']); saldo_y=m['saldo'][11]
        add('8. RESUMO DO MÊS',None,None,'sec')
        add('TOTAL DE DESPESAS (sem investimentos)',m['despesas'],S(m['despesas']),'tot')
        add('TOTAL DE SAÍDAS (despesas + investimentos)',m['saidas'],S(m['saidas']),'tot')
        add('SALDO DO MÊS (sobra para o mês seguinte)',m['saldo'],saldo_y,'bad' if saldo_y<0 else 'tot')
        ratio=lambda a,b:[a[i]/b[i] if b[i] else 0.0 for i in range(12)]
        add('Taxa de poupança (investido ÷ receita)',ratio(m['inv'],m['rec']),S(m['inv'])/rec_y if rec_y else 0.0,'item',fmt=pctf)
        add('% da receita comprometida com despesas',ratio(m['despesas'],m['rec']),S(m['despesas'])/rec_y if rec_y else 0.0,'item',fmt=pctf)
        add('9. CONTROLE: ORÇADO x REALIZADO (quanto ainda sobra em cada bloco)',None,None,'sec')
        for g,p in m['targets'].items():
            ser=self.block_series(m,g); orc=m['orcado'][g]; dif=[orc[i]-ser[i] for i in range(12)]
            add(f'{blk(g)} ({p:g}%) — diferença (orçado − gasto)',dif,S(orc)-S(ser),'bad' if S(orc)-S(ser)<0 else 'item',fmt=br)
            add(f'{blk(g)} ({p:g}%) — % do orçamento usado',ratio(ser,orc),S(ser)/S(orc) if S(orc) else 0.0,'item',fmt=pctf)
        return rows

    def set_cell(self,year,mes,g,cat,valor):
        """Faz o total de uma célula da planilha igualar 'valor'. Devolve mensagem de erro ou None."""
        month=f'{year}-{mes:02d}'; tipo='Receita' if g=='Receitas' else 'Investimento' if g=='Investimentos' else 'Despesa'
        manual=0.0; editable=[]
        for i,t,gr,v,o in self.con.execute("SELECT id,tipo,grupo,valor,origem FROM lancamentos WHERE substr(data,1,7)=? AND categoria=?",(month,cat)).fetchall():
            if eff_group(t,gr,cat)!=g: continue
            if o in ('planilha','grade'): editable.append(i)
            else: manual+=v
        if valor+0.005<manual: return f'Já existem lançamentos manuais somando {money(manual)} nesta célula. Edite ou exclua-os em «Lançamentos».'
        for i in editable: self.con.execute('DELETE FROM lancamentos WHERE id=?',(i,))
        rest=round(valor-manual,2)
        if rest>0: self.con.execute('INSERT INTO lancamentos(data,tipo,categoria,descricao,valor,origem,grupo) VALUES(?,?,?,?,?,?,?)',(f'{month}-01',tipo,cat,'Ajuste na planilha',rest,'grade','' if tipo=='Receita' else g))
        self.con.commit(); return None

    # ------------------------------------------------------------ PAINEL
    def build_dashboard(self):
        self.heading('Visão geral da família','Acompanhe o orçamento, identifique excessos e veja para onde o dinheiro está indo.')
        kp=self.cols(1,1,1,1); self.kpi={}
        for i,(label,color) in enumerate([('Receitas',GREEN),('Despesas',RED),('Investimentos',BLUE),('Saldo final',PURPLE)]):
            box=tk.Frame(kp,bg=WHITE,highlightbackground=BORDER,highlightthickness=1); self.put(kp,box,i)
            tk.Frame(box,bg=color,width=5).pack(side='left',fill='y')
            inner=tk.Frame(box,bg=WHITE); inner.pack(side='left',fill='both',expand=True,padx=14,pady=12)
            tk.Label(inner,text=label.upper(),font=(FONT,8,'bold'),bg=WHITE,fg=MUTED).pack(anchor='w')
            v=tk.Label(inner,text='R$ 0,00',font=(FONT,18,'bold'),bg=WHITE,fg=TEXT); v.pack(anchor='w',pady=(2,0))
            cap=tk.Label(inner,text=' ',font=F_SMALL,bg=WHITE,fg=MUTED); cap.pack(anchor='w'); self.kpi[label]=(v,cap)

        r1=self.cols(3,2)
        c=card(r1,'Receitas, despesas e investimentos','Últimos 6 meses até o mês selecionado'); self.put(r1,c,0)
        self.bar=BarChart(c,height=260); self.bar.pack(fill='both',expand=True,padx=10,pady=(0,12))
        c=card(r1,'Alertas e recomendações','Baseados no orçamento do mês'); self.put(r1,c,1)
        self.alert_frame=tk.Frame(c,bg=WHITE); self.alert_frame.pack(fill='both',expand=True,padx=16,pady=(0,12))

        r2=self.cols(2,3)
        c=card(r2,'Gastos por bloco do orçamento','Distribuição das saídas no mês'); self.put(r2,c,0)
        self.pie_group=DonutChart(c,donut=False,height=250); self.pie_group.pack(fill='both',expand=True,padx=10,pady=(0,12))
        c=card(r2,'Orçamento x realizado','Quanto de cada bloco já foi usado'); self.put(r2,c,1)
        self.budget_bars=HBarList(c); self.budget_bars.pack(fill='both',expand=True,padx=16,pady=(0,12))

        r3=self.cols(1,1)
        c=card(r3,'Despesas por categoria','Principais categorias do mês'); self.put(r3,c,0)
        self.donut_cat=DonutChart(c,donut=True,center_title='Despesas',height=250); self.donut_cat.pack(fill='both',expand=True,padx=10,pady=(0,12))
        c=card(r3,'Origem das receitas','De onde veio o dinheiro no mês'); self.put(r3,c,1)
        self.pie_inc=DonutChart(c,donut=False,height=250); self.pie_inc.pack(fill='both',expand=True,padx=10,pady=(0,12))

    def render_alerts(self,items):
        colors={'danger':RED,'warn':AMBER,'ok':GREEN,'info':BLUE}
        for w in self.alert_frame.winfo_children(): w.destroy()
        for lvl,text in items:
            row=tk.Frame(self.alert_frame,bg=WHITE); row.pack(fill='x',pady=5)
            tk.Label(row,text='●',fg=colors[lvl],bg=WHITE,font=(FONT,10)).pack(side='left',anchor='n')
            tk.Label(row,text=text,bg=WHITE,fg=TEXT2,font=F_BODY,wraplength=300,justify='left').pack(side='left',padx=8)

    def build_alerts(self,month,sums,balance):
        alerts=[]; income=sums['Receita']; spend=self.group_spend(month)
        if balance<0: alerts.append(('danger','O saldo projetado do mês está negativo. Revise os gastos antes de assumir novos compromissos.'))
        if income<=0 and (sums['Despesa']+sums['Investimento'])>0: alerts.append(('warn','Há saídas, mas nenhuma receita registrada. Os limites do orçamento dependem da receita do mês.'))
        for group,pct in self.targets().items():
            budget=income*pct/100; used=spend.get(group,0)
            if budget>0 and used>budget: alerts.append(('danger',f'{blk(group)}: {money(used-budget)} acima do orçamento.'))
            elif budget>0 and used>=budget*.85: alerts.append(('warn',f'{blk(group)}: já usou {used/budget:.0%} do orçamento.'))
        if income>0 and sums['Investimento']==0: alerts.append(('info','Nenhum aporte em investimentos ou reserva de emergência neste mês.'))
        if not alerts: alerts=[('ok','Nenhum excesso identificado com os lançamentos registrados.'),('info','Dica: registre também as pequenas despesas e revise o orçamento no fim do mês.')]
        return alerts

    def refresh_dashboard(self,month,sums):
        if not hasattr(self,'kpi'): return
        opening=self.opening_balance(month); balance=opening+sums['Receita']-sums['Despesa']-sums['Investimento']; income=sums['Receita']
        pct=lambda v:f'{v/income:.0%} da receita' if income>0 else 'sem receita no mês'
        for label,key in [('Receitas','Receita'),('Despesas','Despesa'),('Investimentos','Investimento')]: self.kpi[label][0].configure(text=money(sums[key]))
        self.kpi['Receitas'][1].configure(text='Entradas do mês'); self.kpi['Despesas'][1].configure(text=pct(sums['Despesa'])); self.kpi['Investimentos'][1].configure(text=pct(sums['Investimento']))
        self.kpi['Saldo final'][0].configure(text=money(balance),fg=RED if balance<0 else TEXT); self.kpi['Saldo final'][1].configure(text=f'Saldo anterior: {money(opening)}')
        months=self.last_months(month); data=[self.sums(m) for m in months]
        self.bar.set_data([MONTHS[int(m[5:])-1][:3] for m in months],[('Receitas',GREEN,[d['Receita'] for d in data]),('Despesas',RED,[d['Despesa'] for d in data]),('Investimentos',BLUE,[d['Investimento'] for d in data])])
        spend=self.group_spend(month)
        self.pie_group.set_data(colorize([(blk(g),v) for g,v in spend.items()],LABEL_COLORS))
        self.budget_bars.set_rows([(blk(g),spend.get(g,0),income*p/100,None) for g,p in self.targets().items()])
        self.donut_cat.set_data(colorize(top_n(list(self.by_category(month,'Despesa').items()),7)))
        self.pie_inc.set_data(colorize(top_n(list(self.by_category(month,'Receita').items()),6)))
        self.render_alerts(self.build_alerts(month,sums,balance))

    # ------------------------------------------------------------ PAINEL DO ANO
    def kpi_cards(self,specs):
        row=self.cols(*([1]*len(specs))); out={}
        for i,(label,color) in enumerate(specs):
            box=tk.Frame(row,bg=WHITE,highlightbackground=BORDER,highlightthickness=1); self.put(row,box,i)
            tk.Frame(box,bg=color,width=5).pack(side='left',fill='y'); inner=tk.Frame(box,bg=WHITE); inner.pack(side='left',fill='both',expand=True,padx=14,pady=12)
            tk.Label(inner,text=label.upper(),font=(FONT,8,'bold'),bg=WHITE,fg=MUTED).pack(anchor='w')
            v=tk.Label(inner,text='—',font=(FONT,18,'bold'),bg=WHITE,fg=TEXT); v.pack(anchor='w',pady=(2,0))
            cap=tk.Label(inner,text=' ',font=F_SMALL,bg=WHITE,fg=MUTED); cap.pack(anchor='w'); out[label]=(v,cap)
        return out

    def build_year(self):
        self.heading('Painel do ano','')
        self.year_sub=self.content.winfo_children()[-1]
        self.ykpi=self.kpi_cards([('Receita total do ano',GREEN),('Despesas do ano',RED),('Investimentos do ano',BLUE)])
        self.ykpi.update(self.kpi_cards([('Saldo acumulado do ano',PURPLE),('Taxa de poupança',CYAN),('Gasto médio por mês',AMBER)]))
        c=card(self.content,'Receitas x Despesas x Sobra por mês','O ano completo, mês a mês'); c.pack(fill='x',pady=(0,14))
        self.ybar=BarChart(c,height=270); self.ybar.pack(fill='both',expand=True,padx=10,pady=(0,12))
        r2=self.cols(1,1)
        c=card(r2,'Distribuição dos gastos do ano por grupo','Para onde foi o dinheiro, grupo a grupo'); self.put(r2,c,0)
        self.ypie_group=DonutChart(c,donut=False,height=280); self.ypie_group.pack(fill='both',expand=True,padx=10,pady=(0,12))
        c=card(r2,'Orçado x realizado no ano','Verde: dentro · Amarelo: perto do limite · Vermelho: estourou'); self.put(r2,c,1)
        self.ybars=HBarList(c); self.ybars.pack(fill='both',expand=True,padx=16,pady=(0,12))
        r3=self.cols(1,1)
        c=card(r3,'REAL — para onde foi cada real','Realizado no ano, por bloco'); self.put(r3,c,0)
        self.ypie_real=DonutChart(c,donut=False,height=250); self.ypie_real.pack(fill='both',expand=True,padx=10,pady=(0,12))
        c=card(r3,'IDEAL — seus percentuais-alvo','Como a receita deveria ser dividida (modelo 55 / 20 / 10 / 10 / 5)'); self.put(r3,c,1)
        self.ypie_ideal=DonutChart(c,donut=False,fmt=lambda v:f'{v:g}%',height=250); self.ypie_ideal.pack(fill='both',expand=True,padx=10,pady=(0,12))
        self.ytree=self.table(self.content,('bloco','pct','orcado','real','dif','usado'),('Bloco','% da receita','Orçado no ano','Realizado no ano','Diferença','% usado'),(230,110,150,150,150,110),height=6,right=('pct','orcado','real','dif','usado'))
        self.ytree.master.pack(fill='x',pady=(0,10))

    def refresh_year(self,month,sums):
        if not hasattr(self,'ytree'): return
        year=int(month[:4]); m=self.year_model(year); S=sum; rec=S(m['rec']); desp=S(m['despesas']); inv=S(m['inv'])
        self.year_sub.configure(text=f'Resumo de {year} no modelo da planilha «Orçamento Mensal». Os valores vêm dos seus lançamentos.')
        saldo=m['saldo'][11]; meses=sum(1 for i in range(12) if m['rec'][i]>0 or m['saidas'][i]>0)
        k=self.ykpi
        k['Receita total do ano'][0].configure(text=money(rec)); k['Receita total do ano'][1].configure(text='Soma dos 12 meses')
        k['Despesas do ano'][0].configure(text=money(desp)); k['Despesas do ano'][1].configure(text=f'{desp/rec:.0%} da receita' if rec>0 else 'sem receita no ano')
        k['Investimentos do ano'][0].configure(text=money(inv)); k['Investimentos do ano'][1].configure(text='Total aportado')
        k['Saldo acumulado do ano'][0].configure(text=money(saldo),fg=RED if saldo<0 else TEXT); k['Saldo acumulado do ano'][1].configure(text='Saldo ao fim de dezembro, como na planilha')
        k['Taxa de poupança'][0].configure(text=f'{inv/rec:.1%}' if rec>0 else '—'); k['Taxa de poupança'][1].configure(text='Investido ÷ receita')
        k['Gasto médio por mês'][0].configure(text=money(desp/meses if meses else 0)); k['Gasto médio por mês'][1].configure(text=f'Média de {meses} {"mês" if meses==1 else "meses"} com lançamentos')
        self.ybar.set_data([x[:3] for x in MONTHS],[('Receitas',GREEN,m['rec']),('Despesas',RED,m['despesas']),('Sobra do mês',PURPLE,[max(0.0,v) for v in m['saldo']])])
        items=[]; zeros=[0.0]*12
        for gg,sub,cats in STRUCTURE:
            if sub: items.append((sub.capitalize(),S(S(m['cell'].get((gg,c),zeros)) for c in cats)))
        items.append(('Essencial (outros)',S(m['cell'].get(('Essenciais',OTHER_ROW),zeros))))
        for gg in ('Metas','Gastos livres','Educação'): items.append((blk(gg),S(m['desp_g'][gg])))
        self.ypie_group.set_data(colorize(top_n(items,9)))
        real={g:S(self.block_series(m,g)) for g in BLOCKS}; orc={g:S(m['orcado'][g]) for g in BLOCKS}
        self.ypie_real.set_data(colorize([(blk(g),real[g]) for g in BLOCKS],LABEL_COLORS))
        self.ypie_ideal.set_data(colorize([(blk(g),m['targets'].get(g,0)) for g in BLOCKS],LABEL_COLORS))
        self.ybars.set_rows([(blk(g),real[g],orc[g],None) for g in BLOCKS])
        rows=[(None,(blk(g),f'{m["targets"].get(g,0):g}%',money(orc[g]),money(real[g]),money(orc[g]-real[g]),f'{real[g]/orc[g]:.0%}' if orc[g] else '—'),None) for g in BLOCKS]
        rows.append((None,('Total',f'{S(m["targets"].values()):g}%',money(S(orc.values())),money(S(real.values())),money(S(orc.values())-S(real.values())),f'{S(real.values())/S(orc.values()):.0%}' if S(orc.values()) else '—'),None))
        self.fill(self.ytree,rows)

    # ------------------------------------------------------------ PLANILHA
    def build_grid(self):
        self.heading('Planilha do ano','')
        self.grid_sub=self.content.winfo_children()[-1]
        bar=tk.Frame(self.content,bg=BG); bar.pack(fill='x',pady=(0,8))
        button(bar,'Limpar este mês',lambda:self.clear_entries('month'),'danger').pack(side='left'); button(bar,'Limpar o ano inteiro',lambda:self.clear_entries('year'),'danger').pack(side='left',padx=8)
        tk.Label(bar,text='Para corrigir uma célula: duplo clique, digite o valor certo (deixe em branco para zerar).',font=F_SMALL,bg=BG,fg=MUTED).pack(side='left',padx=12)
        outer=tk.Frame(self.content,bg=WHITE,highlightbackground=BORDER,highlightthickness=1); outer.pack(fill='x',pady=(0,14))
        cols=['cat']+[f'm{i}' for i in range(1,13)]+['ano']
        tree=ttk.Treeview(outer,columns=cols,show='headings',height=26)
        tree.heading('cat',text='CATEGORIA',anchor='w'); tree.column('cat',width=330,anchor='w',stretch=False)
        for i in range(1,13): tree.heading(f'm{i}',text=MONTHS[i-1][:3].upper(),anchor='e'); tree.column(f'm{i}',width=98,anchor='e',stretch=False)
        tree.heading('ano',text='ANO',anchor='e'); tree.column('ano',width=124,anchor='e',stretch=False)
        vs=ttk.Scrollbar(outer,orient='vertical',command=tree.yview); hs=ttk.Scrollbar(outer,orient='horizontal',command=tree.xview)
        tree.configure(yscrollcommand=vs.set,xscrollcommand=hs.set)
        tree.grid(row=0,column=0,sticky='nsew'); vs.grid(row=0,column=1,sticky='ns'); hs.grid(row=1,column=0,sticky='ew'); outer.rowconfigure(0,weight=1); outer.columnconfigure(0,weight=1)
        tree.tag_configure('sec',background='#dbeafe',foreground=TEXT,font=(FONT,10,'bold')); tree.tag_configure('sub',background='#f1f5f9',foreground=TEXT,font=(FONT,10,'bold'))
        tree.tag_configure('subtot',background='#f8fafc',foreground=TEXT,font=(FONT,10,'bold')); tree.tag_configure('tot',background='#e0f2fe',foreground=TEXT,font=(FONT,10,'bold'))
        tree.tag_configure('item',background=WHITE,foreground=TEXT2); tree.tag_configure('bad',background='#fee2e2',foreground='#b91c1c',font=(FONT,10,'bold'))
        tree.bind('<Double-1>',self.on_grid_edit); self.gridtree=tree; self.grid_edit={}; self.grid_year=date.today().year
        tk.Label(self.content,text='Células de totais, percentuais e a divisão orçada são calculadas automaticamente, como na planilha.',font=F_SMALL,bg=BG,fg=MUTED).pack(anchor='w')

    def refresh_grid(self,month,sums):
        if not hasattr(self,'gridtree'): return
        year=int(month[:4]); m=self.year_model(year); rows=self.grid_rows(m,year); tree=self.gridtree; y0=tree.yview()[0]
        self.grid_sub.configure(text=f'Ano de referência: {year}. Dê um duplo clique em uma célula de categoria para digitar o valor do mês. Use ◀ ▶ no topo para trocar de ano/mês.')
        for i in tree.get_children(): tree.delete(i)
        self.grid_edit={}; self.grid_year=year
        for iid,label,cells,ano,tag,edit in rows:
            tree.insert('','end',iid=iid,values=(label,*cells,ano),tags=(tag,))
            if edit is not None: self.grid_edit[iid]=edit
        tree.yview_moveto(y0)

    def on_grid_edit(self,event):
        tree=self.gridtree
        if tree.identify_region(event.x,event.y)!='cell': return
        key=self.grid_edit.get(tree.identify_row(event.y))
        try: idx=int(tree.identify_column(event.x)[1:])-2
        except ValueError: return
        if key is None or not 0<=idx<12: return
        year=self.grid_year; mes=idx+1; month=f'{year}-{mes:02d}'
        if key=='opening':
            if mes!=1: messagebox.showinfo('Sobra do mês anterior','Só janeiro é informado por você. Nos demais meses a sobra é calculada automaticamente a partir do saldo do mês anterior.'); return
            cur=self.opening_balance(month); raw=simpledialog.askstring('Sobra do mês anterior',f'Quanto a família já tinha guardado antes de janeiro/{year}? (R$)',initialvalue=f'{cur:.2f}'.replace('.',','),parent=self)
            if raw is None: return
            try:
                val=parse_amount(raw)
                if val<0: raise ValueError
            except ValueError: messagebox.showerror('Valor inválido','Digite um valor válido. Exemplo: 1.250,00'); return
            self.con.execute('INSERT OR REPLACE INTO saldos_iniciais(mes,valor) VALUES(?,?)',(month,val)); self.con.commit(); self.refresh_page(); return
        g,cat=key; cur=self.year_model(year)['cell'].get((g,cat),[0.0]*12)[idx]
        raw=simpledialog.askstring(cat,f'{MONTHS[idx]} de {year}\nTotal do mês para «{cat}» (R$):\nDeixe em branco para zerar.',initialvalue=f'{cur:.2f}'.replace('.',',') if cur else '',parent=self)
        if raw is None: return
        try:
            val=0.0 if raw.strip()=='' else parse_amount(raw)
            if val<0: raise ValueError
        except ValueError: messagebox.showerror('Valor inválido','Digite um valor válido. Exemplo: 125,50.'); return
        err=self.set_cell(year,mes,g,cat,val)
        if err: messagebox.showwarning('Não foi possível alterar',err)
        self.refresh_page()

    # ------------------------------------------------------------ LANÇAMENTOS
    def build_mov(self):
        self.heading('Lançamentos','Registre receitas, despesas e aportes. Errou algo? Selecione a linha e clique em «Editar selecionado» (ou dê duplo clique), corrija e salve.')
        self.editing_id=None
        form=card(self.content); form.pack(fill='x',pady=(0,14))
        self.f_data=tk.StringVar(value=date.today().isoformat()); self.f_tipo=tk.StringVar(value='Despesa'); self.f_bloco=tk.StringVar(); self.f_cat=tk.StringVar(); self.f_desc=tk.StringVar(); self.f_valor=tk.StringVar()
        self.labeled_entry(form,'Data',self.f_data,11,0)
        self.labeled_entry(form,'Tipo',self.f_tipo,11,1,widget='combo',values=['Receita','Despesa','Investimento'])
        self.cb_bloco=self.labeled_entry(form,'Bloco do orçamento',self.f_bloco,22,2,widget='combo',values=[])
        self.cb_cat=self.labeled_entry(form,'Categoria',self.f_cat,22,3,widget='combo',values=[])
        self.labeled_entry(form,'Descrição',self.f_desc,22,4)
        e=self.labeled_entry(form,'Valor (R$)',self.f_valor,12,5); e.bind('<Return>',lambda ev:self.add_entry())
        self.f_tipo.trace_add('write',lambda *a:self._sync_tipo()); self.f_bloco.trace_add('write',lambda *a:self._sync_bloco()); self._sync_tipo()
        self.entry_btn=button(form,'＋  Adicionar lançamento',self.add_entry); self.entry_btn.grid(row=2,column=0,columnspan=2,sticky='w',padx=10,pady=(0,14))
        button(form,'Limpar formulário',self.reset_form,'ghost').grid(row=2,column=2,sticky='w',padx=10,pady=(0,14))
        self.form_note=tk.Label(form,text='',font=(FONT,9,'bold'),bg=WHITE,fg=AMBER); self.form_note.grid(row=2,column=3,columnspan=3,sticky='w',padx=10,pady=(0,14))
        self.mov_summary=tk.Label(self.content,text='',font=(FONT,10,'bold'),bg=BG,fg=TEXT2); self.mov_summary.pack(anchor='w',pady=(0,8))
        self.entries=self.table(self.content,('data','tipo','bloco','categoria','descricao','valor'),('Data','Tipo','Bloco','Categoria','Descrição','Valor'),(105,105,170,210,290,120),height=15,right=('valor',))
        self.entries.master.pack(fill='x'); self.entries.bind('<Double-1>',lambda e:self.edit_entry())
        row=tk.Frame(self.content,bg=BG); row.pack(fill='x',pady=12)
        button(row,'✎  Editar selecionado',self.edit_entry,'ghost').pack(side='left')
        button(row,'Excluir selecionados',self.delete_entry,'danger').pack(side='right')
        button(row,'Limpar lançamentos do mês',lambda:self.clear_entries('month'),'danger').pack(side='right',padx=8)

    def _sync_tipo(self):
        tipo=self.f_tipo.get()
        blocos={'Receita':['Receitas'],'Despesa':[BLOCK_LONG[g] for g in ('Essenciais','Metas','Gastos livres','Educação')],'Investimento':[BLOCK_LONG['Investimentos']]}.get(tipo,[])
        self.cb_bloco.configure(values=blocos)
        if blocos and self.f_bloco.get() not in blocos: self.f_bloco.set(blocos[0])
        else: self._sync_bloco()

    def _sync_bloco(self):
        label=self.f_bloco.get(); key=LONG_TO_KEY.get(label)
        cats=INCOME_CATS if label=='Receitas' else [c for g,_,cs in STRUCTURE if g==key for c in cs]
        self.cb_cat.configure(values=cats)
        if cats and self.f_cat.get() not in cats: self.f_cat.set(cats[0])

    def refresh_mov(self,month,sums):
        if not hasattr(self,'entries'): return
        rows=[(r[0],(r[1],r[2],'—' if r[2]=='Receita' else blk(eff_group(r[2],r[7],r[3])),r[3],r[4],money(r[5])),r[2]) for r in self.month_rows(month)]
        self.fill(self.entries,rows)
        self.mov_summary.configure(text=f'Receitas {money(sums["Receita"])}   ·   Despesas {money(sums["Despesa"])}   ·   Investimentos {money(sums["Investimento"])}')

    # ------------------------------------------------------------ ORÇAMENTO
    def build_budget(self):
        self.heading('Orçamento mensal','A distribuição é uma referência inicial. Ajuste-a à renda e às necessidades reais da família.')
        self.targetvars={k:tk.StringVar(value=self.pref(f'alvo_{k}',v)) for k,v in TARGETS_DEFAULT.items()}
        r1=self.cols(3,2)
        c=card(r1,'Orçado x realizado','Base do rateio: receita do mês, sem o saldo transportado'); self.put(r1,c,0)
        self.targettree=self.table(c,('grupo','percentual','valor','realizado','diferenca'),('Bloco','% alvo','Orçado','Realizado','Disponível / excedido'),(150,80,120,120,160),height=6,right=('percentual','valor','realizado','diferenca'))
        self.targettree.master.pack(fill='x',padx=14,pady=(0,14))
        c=card(r1,'Distribuição-alvo','Como a receita deve ser dividida'); self.put(r1,c,1)
        self.pie_targets=DonutChart(c,donut=False,fmt=lambda v:f'{v:g}%',height=230); self.pie_targets.pack(fill='both',expand=True,padx=10,pady=(0,12))
        r2=self.cols(3,2)
        c=card(r2,'Consumo do orçamento','Verde: dentro do limite · Amarelo: acima de 85% · Vermelho: estourou'); self.put(r2,c,0)
        self.budget_bars2=HBarList(c); self.budget_bars2.pack(fill='both',expand=True,padx=16,pady=(0,12))
        c=card(r2,'Editar percentuais','Os cinco blocos devem somar 100%'); self.put(r2,c,1)
        form=tk.Frame(c,bg=WHITE); form.pack(fill='x',padx=16)
        for i,(k,var) in enumerate(self.targetvars.items()):
            tk.Frame(form,bg=GROUP_COLORS[k],width=10,height=10).grid(row=i,column=0,padx=(0,8),pady=6)
            tk.Label(form,text=blk(k),font=F_BODY,bg=WHITE,fg=TEXT2).grid(row=i,column=1,sticky='w'); ttk.Entry(form,textvariable=var,width=8,justify='right').grid(row=i,column=2,padx=(14,4),pady=6); tk.Label(form,text='%',bg=WHITE,fg=MUTED).grid(row=i,column=3)
            var.trace_add('write',lambda *a:self.refresh_budget(self.get_month(),None))
        form.columnconfigure(1,weight=1)
        self.sum_lbl=tk.Label(c,text='',font=(FONT,10,'bold'),bg=WHITE); self.sum_lbl.pack(anchor='w',padx=16,pady=(8,4))
        brow=tk.Frame(c,bg=WHITE); brow.pack(anchor='w',padx=16,pady=(4,16)); button(brow,'Salvar percentuais',self.save_targets).pack(side='left'); button(brow,'Restaurar padrão',self.restore_targets,'ghost').pack(side='left',padx=8)

    def refresh_budget(self,month,sums):
        if self.current_page!='Orçamento' or not hasattr(self,'targettree'): return
        sums=sums or self.sums(month); targets={k:safe_amount(v.get()) for k,v in self.targetvars.items()}; spend=self.group_spend(month); income=sums['Receita']
        self.fill(self.targettree,[(None,(blk(g),f'{p:g}%',money(income*p/100),money(spend.get(g,0)),money(income*p/100-spend.get(g,0))),None) for g,p in targets.items()])
        self.pie_targets.set_data(colorize([(blk(g),p) for g,p in targets.items()],LABEL_COLORS))
        self.budget_bars2.set_rows([(blk(g),spend.get(g,0),income*p/100,None) for g,p in targets.items()])
        total=sum(targets.values()); ok=abs(total-100)<=0.01
        self.sum_lbl.configure(text=f'Soma atual: {total:g}%'+('  ✓' if ok else '  (precisa ser 100%)'),fg=GREEN if ok else RED)

    def save_targets(self):
        try:
            vals={k:parse_amount(v.get()) for k,v in self.targetvars.items()}
            if any(v<0 for v in vals.values()) or abs(sum(vals.values())-100)>0.01: raise ValueError
            for k,v in vals.items(): self.con.execute('INSERT OR REPLACE INTO preferencias(chave,valor) VALUES(?,?)',(f'alvo_{k}',str(v)))
            self.con.commit(); messagebox.showinfo('Orçamento','Percentuais salvos.'); self.refresh_page()
        except Exception: messagebox.showerror('Percentuais inválidos','Os percentuais devem ser positivos e somar 100%.')

    # ------------------------------------------------------------ METAS
    def build_goals(self):
        self.heading('Metas financeiras','Defina objetivos, acompanhe o progresso e saiba quanto ainda falta. Para corrigir uma meta, selecione-a na tabela e clique em «Editar meta».')
        self.editing_goal=None
        form=card(self.content); form.pack(fill='x',pady=(0,14))
        self.g_nome=tk.StringVar(); self.g_alvo=tk.StringVar(); self.g_atual=tk.StringVar(value='0'); self.g_prazo=tk.StringVar()
        for i,(lab,var) in enumerate([('Nome da meta',self.g_nome),('Valor alvo (R$)',self.g_alvo),('Já guardado (R$)',self.g_atual),('Prazo (AAAA-MM)',self.g_prazo)]): self.labeled_entry(form,lab,var,22,i)
        self.goal_btn=button(form,'＋  Adicionar meta',self.add_goal); self.goal_btn.grid(row=2,column=0,padx=10,pady=(0,14),sticky='w')
        button(form,'Limpar formulário',self.reset_goal_form,'ghost').grid(row=2,column=1,padx=10,pady=(0,14),sticky='w')
        self.goal_note=tk.Label(form,text='',font=(FONT,9,'bold'),bg=WHITE,fg=AMBER); self.goal_note.grid(row=2,column=2,columnspan=2,sticky='w',padx=10,pady=(0,14))
        r1=self.cols(3,2)
        c=card(r1,'Progresso das metas','Quanto já foi guardado em relação ao alvo'); self.put(r1,c,0)
        self.goal_bars=HBarList(c); self.goal_bars.pack(fill='both',expand=True,padx=16,pady=(0,12))
        c=card(r1,'Quanto já guardei por meta','Distribuição do valor acumulado'); self.put(r1,c,1)
        self.pie_goals=DonutChart(c,donut=True,center_title='Guardado',height=230); self.pie_goals.pack(fill='both',expand=True,padx=10,pady=(0,12))
        self.goalview=self.table(self.content,('nome','alvo','atual','falta','prazo','progresso','aporte'),('Meta','Valor alvo','Já guardado','Falta','Prazo','Progresso','Aporte mensal sugerido'),(230,120,120,120,90,90,170),height=8,right=('alvo','atual','falta','progresso','aporte'))
        self.goalview.master.pack(fill='x'); self.goalview.bind('<Double-1>',lambda e:self.edit_goal())
        row=tk.Frame(self.content,bg=BG); row.pack(fill='x',pady=12)
        button(row,'✎  Editar meta',self.edit_goal,'ghost').pack(side='left'); button(row,'Atualizar valor guardado',self.update_goal,'ghost').pack(side='left',padx=8)
        button(row,'Excluir selecionadas',self.delete_goal,'danger').pack(side='right'); button(row,'Limpar todas as metas',self.clear_goals,'danger').pack(side='right',padx=8)

    def refresh_goals(self,month,sums):
        if not hasattr(self,'goalview'): return
        rows=self.con.execute('SELECT id,nome,valor_alvo,valor_atual,prazo FROM metas ORDER BY id DESC').fetchall(); table=[]; bars=[]
        for i,nome,alvo,atual,prazo in rows:
            progress=min(100,atual/alvo*100) if alvo else 0; falta=max(0,alvo-atual); left=months_until(prazo) if prazo else None
            if falta<=0: aporte='Meta concluída'
            elif left is None: aporte='—'
            elif left<0: aporte='Prazo vencido'
            else: aporte=money(falta/max(left,1))+'/mês'
            table.append((i,(nome,money(alvo),money(atual),money(falta),prazo or '—',f'{progress:.0f}%',aporte),None))
            bars.append((nome,atual,alvo,GREEN if progress>=100 else PRIMARY))
        self.fill(self.goalview,table); self.goal_bars.set_rows(bars)
        self.pie_goals.set_data(colorize([(n,a) for n,a,_,_ in bars]))

    def add_goal(self):
        try:
            nome=self.g_nome.get().strip(); alvo=parse_amount(self.g_alvo.get()); atual=parse_amount(self.g_atual.get()); prazo=self.g_prazo.get().strip()
            if not nome or alvo<=0 or atual<0: raise ValueError
            if prazo: datetime.strptime(prazo,'%Y-%m')
        except Exception:
            messagebox.showerror('Dados inválidos','Informe nome, alvo maior que zero, valor guardado válido e prazo AAAA-MM opcional.'); return
        self.store_goal(nome,alvo,atual,prazo,self.editing_goal); self.reset_goal_form(); self.refresh_page()

    def store_goal(self,nome,alvo,atual,prazo,goal_id=None):
        if goal_id: self.con.execute('UPDATE metas SET nome=?,valor_alvo=?,valor_atual=?,prazo=? WHERE id=?',(nome,alvo,atual,prazo,goal_id))
        else: self.con.execute('INSERT INTO metas(nome,valor_alvo,valor_atual,prazo) VALUES(?,?,?,?)',(nome,alvo,atual,prazo))
        self.con.commit()

    def reset_goal_form(self):
        self.editing_goal=None; self.g_nome.set(''); self.g_alvo.set(''); self.g_atual.set('0'); self.g_prazo.set('')
        self.goal_btn.configure(text='＋  Adicionar meta'); self.goal_note.configure(text='')

    def edit_goal(self):
        sel=self.goalview.selection()
        if len(sel)!=1: messagebox.showinfo('Editar meta','Selecione uma única meta na tabela.'); return
        r=self.con.execute('SELECT id,nome,valor_alvo,valor_atual,prazo FROM metas WHERE id=?',(int(sel[0]),)).fetchone()
        if not r: return
        i,nome,alvo,atual,prazo=r; fm=lambda v:f'{v:.2f}'.replace('.',',')
        self.g_nome.set(nome); self.g_alvo.set(fm(alvo)); self.g_atual.set(fm(atual)); self.g_prazo.set(prazo or ''); self.editing_goal=i
        self.goal_btn.configure(text='✓  Salvar alterações'); self.goal_note.configure(text='Editando uma meta: altere os campos acima e salve.')

    def update_goal(self):
        sel=self.goalview.selection()
        if not sel: messagebox.showinfo('Atualizar meta','Selecione uma meta na tabela.'); return
        raw=simpledialog.askstring('Atualizar valor guardado','Novo valor guardado (R$):',parent=self)
        if raw is None: return
        try:
            val=parse_amount(raw)
            if val<0: raise ValueError
        except ValueError: messagebox.showerror('Valor inválido','Digite um valor válido. Exemplo: 1.250,00'); return
        self.con.execute('UPDATE metas SET valor_atual=? WHERE id=?',(val,int(sel[0]))); self.con.commit(); self.refresh_page()

    def delete_goal(self):
        sel=self.goalview.selection()
        if not sel: messagebox.showinfo('Excluir metas','Selecione uma ou mais metas na tabela.'); return
        if messagebox.askyesno('Confirmar exclusão',f'Excluir {len(sel)} meta(s) selecionada(s)?'):
            ids=[int(i) for i in sel]
            for i in ids: self.con.execute('DELETE FROM metas WHERE id=?',(i,))
            self.con.commit()
            if self.editing_goal in ids: self.reset_goal_form()
            self.refresh_page()

    # ------------------------------------------------------------ INVESTIMENTOS
    def build_investments(self):
        self.heading('Investimentos e patrimônio','Controle os aportes por classe. Este painel não recomenda ativos nem promete rentabilidade.')
        kp=self.cols(1,1,1)
        self.inv_kpi={}
        for i,(label,color) in enumerate([('Aportes no mês',GREEN),('Aportes no ano',BLUE),('% da receita do mês',PURPLE)]):
            box=tk.Frame(kp,bg=WHITE,highlightbackground=BORDER,highlightthickness=1); self.put(kp,box,i)
            tk.Frame(box,bg=color,width=5).pack(side='left',fill='y'); inner=tk.Frame(box,bg=WHITE); inner.pack(side='left',fill='both',expand=True,padx=14,pady=12)
            tk.Label(inner,text=label.upper(),font=(FONT,8,'bold'),bg=WHITE,fg=MUTED).pack(anchor='w'); v=tk.Label(inner,text='—',font=(FONT,18,'bold'),bg=WHITE,fg=TEXT); v.pack(anchor='w',pady=(2,0)); self.inv_kpi[label]=v
        r1=self.cols(2,3)
        c=card(r1,'Aportes por classe','Participação no mês selecionado'); self.put(r1,c,0)
        self.pie_inv=DonutChart(c,donut=False,height=250); self.pie_inv.pack(fill='both',expand=True,padx=10,pady=(0,12))
        c=card(r1,'Evolução dos aportes','Últimos 6 meses'); self.put(r1,c,1)
        self.bar_inv=BarChart(c,height=250); self.bar_inv.pack(fill='both',expand=True,padx=10,pady=(0,12))
        self.inv_tree=self.table(self.content,('categoria','total','percentual'),('Categoria','Aportes no mês','Participação'),(360,220,180),height=8,right=('total','percentual'))
        self.inv_tree.master.pack(fill='x')
        tk.Label(self.content,text='Antes de investir, avalie risco, liquidez, custos, impostos, prazo e adequação ao seu perfil.',font=F_BODY,bg=BG,fg=MUTED).pack(anchor='w',pady=14)

    def refresh_investments(self,month,sums):
        if not hasattr(self,'inv_tree'): return
        cats=self.by_category(month,'Investimento'); total=sums['Investimento']; year=month[:4]
        ano=self.con.execute("SELECT COALESCE(SUM(valor),0) FROM lancamentos WHERE tipo='Investimento' AND substr(data,1,4)=?",(year,)).fetchone()[0]
        self.inv_kpi['Aportes no mês'].configure(text=money(total)); self.inv_kpi['Aportes no ano'].configure(text=money(ano))
        self.inv_kpi['% da receita do mês'].configure(text=f'{total/sums["Receita"]:.1%}' if sums['Receita']>0 else '—')
        self.fill(self.inv_tree,[(None,(cat,money(val),f'{val/total:.1%}' if total else '0%'),None) for cat,val in sorted(cats.items(),key=lambda x:x[1],reverse=True)])
        self.pie_inv.set_data(colorize(top_n(list(cats.items()),7)))
        months=self.last_months(month); self.bar_inv.set_data([MONTHS[int(m[5:])-1][:3] for m in months],[('Aportes',GREEN,[self.sums(m)['Investimento'] for m in months])])

    # ------------------------------------------------------------ RELATÓRIOS
    def build_reports(self):
        self.heading('Relatórios e segurança','Veja o ano completo, exporte os dados, faça cópias de segurança e corrija ou limpe informações preenchidas errado.')
        c=card(self.content,'Resultado do ano','Receitas, despesas e investimentos mês a mês'); c.pack(fill='x',pady=(0,14))
        self.bar_year=BarChart(c,height=270); self.bar_year.pack(fill='both',expand=True,padx=10,pady=(0,12))
        panel=card(self.content,'Exportar e proteger seus dados'); panel.pack(fill='x',pady=(0,14))
        row=tk.Frame(panel,bg=WHITE); row.pack(fill='x',padx=16,pady=(0,16))
        for text,cmd,kind in [('Exportar Excel (.xlsx)',self.export_excel,'primary'),('Exportar CSV',self.export_csv,'ghost'),('Criar backup',self.backup_db,'ghost'),('Restaurar backup',self.restore_db,'danger')]: button(row,text,cmd,kind).pack(side='left',padx=(0,10))
        fix=card(self.content,'Corrigir ou limpar informações','Preencheu algo errado? Apague só o que precisa, sem perder o resto. Para corrigir um item, use «Editar» nas telas Lançamentos e Metas, ou dê duplo clique na Planilha.'); fix.pack(fill='x',pady=(0,14))
        acts=[('Limpar lançamentos do mês selecionado',lambda:self.clear_entries('month')),('Limpar lançamentos do ano selecionado',lambda:self.clear_entries('year')),
              ('Limpar valores importados do Excel',lambda:self.clear_entries('imported')),('Limpar todas as metas',self.clear_goals),('Restaurar percentuais 55 / 20 / 10 / 10 / 5',self.restore_targets)]
        for k in range(0,len(acts),2):
            row=tk.Frame(fix,bg=WHITE); row.pack(fill='x',padx=16,pady=(0,8))
            for text,cmd in acts[k:k+2]: button(row,text,cmd,'ghost').pack(side='left',padx=(0,10))
        tk.Frame(fix,height=8,bg=WHITE).pack()
        danger=card(self.content,'Restaurar o programa do zero','Apaga TODOS os lançamentos, metas, sobra inicial e percentuais e deixa o programa como novo. Uma cópia de segurança é criada automaticamente antes e a ação pede confirmação.'); danger.pack(fill='x',pady=(0,14))
        button(danger,'⟲  Restaurar do zero',self.reset_all,'danger').pack(anchor='w',padx=16,pady=(0,16))
        tk.Label(self.content,text=f'Banco de dados atual: {DB}',bg=BG,fg=MUTED,font=F_SMALL,wraplength=900,justify='left').pack(anchor='w',pady=6)

    def refresh_reports(self,month,sums):
        if not hasattr(self,'bar_year'): return
        year=month[:4]; months=[f'{year}-{m:02d}' for m in range(1,13)]; data=[self.sums(m) for m in months]
        self.bar_year.set_data([x[:3] for x in MONTHS],[('Receitas',GREEN,[d['Receita'] for d in data]),('Despesas',RED,[d['Despesa'] for d in data]),('Investimentos',BLUE,[d['Investimento'] for d in data])])

    # ------------------------------------------------------------ APRENDER
    def build_lessons(self):
        self.heading('Aprender educação financeira','O método da planilha, simuladores para testar ideias e lições rápidas. Conteúdo educativo, sem substituir orientação profissional individual.')
        c=card(self.content,'Como usar este programa','O mesmo passo a passo da planilha «Organização Financeira Para Família»'); c.pack(fill='x',pady=(0,14))
        steps=['Em «Planilha» ou «Lançamentos», registre a receita do mês (salários e extras).',
               'Em janeiro, informe quanto a família já tinha guardado na linha «Sobra do mês anterior» da Planilha (duplo clique). Nos outros meses ela é automática.',
               'Confira os percentuais em «Orçamento». Pode mudar, mas a soma precisa dar 100%.',
               'Lance os gastos de cada bloco, mês a mês: pela Planilha ou, no dia a dia, em «Lançamentos».',
               'Acompanhe o «Painel»: ele mostra quanto ainda sobra em cada bloco e avisa quando o gasto passa do orçado.',
               'No fim do ano, veja o «Painel do ano» e compare o REAL com o IDEAL.']
        for n,txt in enumerate(steps,1):
            row=tk.Frame(c,bg=WHITE); row.pack(fill='x',padx=16,pady=3)
            tk.Label(row,text=str(n),font=(FONT,10,'bold'),bg=PRIMARY,fg='white',width=3).pack(side='left',anchor='n')
            tk.Label(row,text=txt,font=F_BODY,bg=WHITE,fg=TEXT2,wraplength=880,justify='left').pack(side='left',padx=10)
        tk.Frame(c,height=10,bg=WHITE).pack()
        c=card(self.content,'O método 55 / 20 / 10 / 10 / 5','Divisão sugerida da receita do mês. É um ponto de partida: ajuste em «Orçamento» conforme a realidade da família.'); c.pack(fill='x',pady=(0,14))
        t=self.targets()
        for g in ('Essenciais','Metas','Investimentos','Gastos livres','Educação'):
            row=tk.Frame(c,bg=WHITE); row.pack(fill='x',padx=16,pady=4)
            tk.Frame(row,bg=GROUP_COLORS[g],width=12,height=12).pack(side='left',anchor='n',pady=4)
            tk.Label(row,text=f'{BLOCK_LONG[g]} · {t.get(g,0):g}%',font=(FONT,10,'bold'),bg=WHITE,fg=TEXT,width=34,anchor='w').pack(side='left',padx=10,anchor='n')
            tk.Label(row,text=BLOCK_HELP[g],font=F_BODY,bg=WHITE,fg=TEXT2,wraplength=560,justify='left').pack(side='left',anchor='n')
        tk.Frame(c,height=10,bg=WHITE).pack()

        def money_txt(v): return f'{v:.2f}'.replace('.',',')
        month=self.get_month(); inc=self.sums(month)['Receita']
        vals=[v for v in (self.group_spend(mm).get('Essenciais',0) for mm in self.last_months(month,3)) if v>0]; avg=sum(vals)/len(vals) if vals else 0
        self.s_renda=tk.StringVar(value=money_txt(inc if inc>0 else 5000)); self.r_desp=tk.StringVar(value=money_txt(avg) if avg else ''); self.r_meses=tk.StringVar(value='6'); self.r_aporte=tk.StringVar(value='')
        self.j_aporte=tk.StringVar(value='300'); self.j_taxa=tk.StringVar(value='8'); self.j_anos=tk.StringVar(value='10'); self.j_inicial=tk.StringVar(value='0')
        def field(parent,text,var,w=12):
            r=tk.Frame(parent,bg=WHITE); r.pack(fill='x',padx=16,pady=3); tk.Label(r,text=text,font=F_BODY,bg=WHITE,fg=TEXT2).pack(side='left'); ttk.Entry(r,textvariable=var,width=w,justify='right').pack(side='right')

        sim=self.cols(1,1)
        c=card(sim,'Simulador: divisão da sua renda','Digite a renda do mês e veja quanto vai para cada bloco'); self.put(sim,c,0)
        field(c,'Renda do mês (R$)',self.s_renda,14)
        self.s_pie=DonutChart(c,donut=False,height=210); self.s_pie.pack(fill='x',padx=10,pady=(6,0))
        self.s_list=tk.Frame(c,bg=WHITE); self.s_list.pack(fill='x',padx=16,pady=(4,14))
        c=card(sim,'Simulador: reserva de emergência','Quanto guardar e em quanto tempo chegar lá'); self.put(sim,c,1)
        field(c,'Despesas essenciais por mês (R$)',self.r_desp,14); field(c,'Meses de proteção desejados',self.r_meses,14); field(c,'Quanto consigo guardar por mês (R$)',self.r_aporte,14)
        self.r_out=tk.Label(c,text='',font=(FONT,11,'bold'),bg=WHITE,fg=TEXT,wraplength=420,justify='left'); self.r_out.pack(anchor='w',padx=16,pady=(14,6))
        tk.Label(c,text='Referência comum: de 3 a 6 meses de despesas essenciais quando a renda é estável; considere mais meses se a renda for variável. A reserva é uma das primeiras prioridades antes de investir em ativos de risco.',font=F_SMALL,bg=WHITE,fg=MUTED,wraplength=420,justify='left').pack(anchor='w',padx=16,pady=(0,14))

        c=card(self.content,'Simulador: juros compostos','Veja o efeito do tempo e dos aportes regulares. Simulação ilustrativa, com taxa constante, sem impostos nem inflação: não é promessa de rentabilidade.'); c.pack(fill='x',pady=(0,14))
        fr=tk.Frame(c,bg=WHITE); fr.pack(fill='x',padx=16)
        for text,var in (('Aporte mensal (R$)',self.j_aporte),('Taxa ao ano (%)',self.j_taxa),('Prazo (anos)',self.j_anos),('Valor inicial (R$)',self.j_inicial)):
            col=tk.Frame(fr,bg=WHITE); col.pack(side='left',padx=(0,18)); tk.Label(col,text=text,font=F_LABEL,bg=WHITE,fg=MUTED).pack(anchor='w'); ttk.Entry(col,textvariable=var,width=14,justify='right').pack(pady=(2,8))
        self.j_out=tk.Label(c,text='',font=(FONT,11,'bold'),bg=WHITE,fg=TEXT,justify='left'); self.j_out.pack(anchor='w',padx=16,pady=(0,4))
        self.j_bar=BarChart(c,height=250); self.j_bar.pack(fill='both',expand=True,padx=10,pady=(0,12))
        for v in (self.s_renda,self.r_desp,self.r_meses,self.r_aporte,self.j_aporte,self.j_taxa,self.j_anos,self.j_inicial): v.trace_add('write',lambda *a:self.update_sims())

        tk.Label(self.content,text='Lições rápidas',font=F_H3,bg=BG,fg=TEXT).pack(anchor='w',pady=(6,4))
        for title,body in LESSONS:
            box=card(self.content); box.pack(fill='x',pady=5)
            num_,_,name=title.partition('. ')
            tk.Label(box,text=num_,font=(FONT,14,'bold'),bg=PRIMARY,fg='white',width=3,pady=6).pack(side='left',anchor='n',padx=16,pady=16)
            txt=tk.Frame(box,bg=WHITE); txt.pack(side='left',fill='x',expand=True,pady=14,padx=(0,16))
            tk.Label(txt,text=name,font=F_H3,bg=WHITE,fg=TEXT).pack(anchor='w',pady=(0,4)); tk.Label(txt,text=body,font=F_BODY,bg=WHITE,fg=TEXT2,wraplength=820,justify='left').pack(anchor='w')
        self.update_sims()

    def update_sims(self):
        if self.current_page!='Aprender' or not hasattr(self,'s_pie'): return
        parts=split_income(safe_amount(self.s_renda.get()),self.targets())
        self.s_pie.set_data(colorize([(blk(g),v) for g,_,v in parts],LABEL_COLORS))
        for w in self.s_list.winfo_children(): w.destroy()
        for g,p,v in parts:
            row=tk.Frame(self.s_list,bg=WHITE); row.pack(fill='x',pady=1)
            tk.Label(row,text=f'{blk(g)} ({p:g}%)',bg=WHITE,fg=TEXT2,font=F_BODY).pack(side='left'); tk.Label(row,text=money(v),bg=WHITE,fg=TEXT,font=(FONT,10,'bold')).pack(side='right')
        meta=reserve_target(safe_amount(self.r_desp.get()),safe_amount(self.r_meses.get())); aporte=safe_amount(self.r_aporte.get())
        if meta<=0: txt='Informe suas despesas essenciais mensais e os meses de proteção.'
        else:
            txt=f'Meta da reserva: {money(meta)}'; n=months_to_goal(meta,aporte)
            if n is not None: txt+=f'\nGuardando {money(aporte)} por mês, você chega lá em {duration_text(n)}.'
            else: txt+='\nInforme quanto consegue guardar por mês para estimar o prazo.'
        self.r_out.configure(text=txt)
        anos=max(0,min(50,int(safe_amount(self.j_anos.get())))); proj=compound_projection(safe_amount(self.j_aporte.get()),safe_amount(self.j_taxa.get()),anos,safe_amount(self.j_inicial.get()))
        if proj:
            a,inv,sal=proj[-1]; self.j_out.configure(text=f'Em {a} {"ano" if a==1 else "anos"}: investido {money(inv)}  ·  saldo estimado {money(sal)}  ·  rendimento {money(sal-inv)}')
        else: self.j_out.configure(text='Informe um prazo de pelo menos 1 ano.')
        self.j_bar.set_data([f'{a}a' for a,_,_ in proj],[('Total investido',BLUE,[i for _,i,_ in proj]),('Saldo estimado',GREEN,[b for _,_,b in proj])])

    # ------------------------------------------------------------ atualização
    def refresh_page(self):
        if not hasattr(self,'content'): return
        month=self.get_month(); sums=self.sums(month); self.month_lbl.configure(text=month_label(month))
        if self.top_month.get()!=month: self.top_month.set(month)
        if hasattr(self,'month_combo'): self.month_combo.configure(values=self.available_months())
        {'Painel':self.refresh_dashboard,'Painel do ano':self.refresh_year,'Planilha':self.refresh_grid,'Lançamentos':self.refresh_mov,'Orçamento':self.refresh_budget,'Metas':self.refresh_goals,'Investimentos':self.refresh_investments,'Relatórios':self.refresh_reports}.get(self.current_page,lambda m,s:None)(month,sums)

    # ------------------------------------------------------------ ações
    def add_entry(self):
        try:
            d=self.f_data.get().strip(); date.fromisoformat(d); tipo=self.f_tipo.get(); cat=self.f_cat.get().strip(); desc=self.f_desc.get().strip(); val=parse_amount(self.f_valor.get())
            if tipo not in ('Receita','Despesa','Investimento') or not cat or val<=0: raise ValueError
            grupo='' if tipo=='Receita' else 'Investimentos' if tipo=='Investimento' else LONG_TO_KEY.get(self.f_bloco.get(),self.category_group(cat) or 'Essenciais')
        except Exception:
            messagebox.showerror('Dados inválidos','Confira a data (AAAA-MM-DD), a categoria e o valor (maior que zero). Exemplo de valor: 125,50.'); return
        self.store_entry(d,tipo,cat,desc,val,grupo,self.editing_id)
        self.reset_form(); self.top_month.set(d[:7]); self.refresh_page()

    def store_entry(self,d,tipo,cat,desc,val,grupo,entry_id=None):
        """Insere um lançamento ou, com entry_id, reescreve o existente (passa a contar como lançamento manual)."""
        if entry_id: self.con.execute("UPDATE lancamentos SET data=?,tipo=?,categoria=?,descricao=?,valor=?,origem='manual',grupo=? WHERE id=?",(d,tipo,cat,desc,val,grupo,entry_id))
        else: self.con.execute('INSERT INTO lancamentos(data,tipo,categoria,descricao,valor,origem,grupo) VALUES(?,?,?,?,?,?,?)',(d,tipo,cat,desc,val,'manual',grupo))
        self.con.commit()

    def reset_form(self):
        self.editing_id=None; self.f_desc.set(''); self.f_valor.set(''); self.f_data.set(date.today().isoformat())
        self.entry_btn.configure(text='＋  Adicionar lançamento'); self.form_note.configure(text='')

    def edit_entry(self):
        sel=self.entries.selection()
        if len(sel)!=1: messagebox.showinfo('Editar lançamento','Selecione um único lançamento na tabela.'); return
        r=self.con.execute('SELECT id,data,tipo,categoria,descricao,valor,grupo FROM lancamentos WHERE id=?',(int(sel[0]),)).fetchone()
        if not r: return
        i,d,tipo,cat,desc,val,grupo=r; g=eff_group(tipo,grupo,cat)
        self.f_data.set(d); self.f_tipo.set(tipo); self.f_bloco.set('Receitas' if g=='Receitas' else BLOCK_LONG.get(g,'Essencial')); self.f_cat.set(cat)
        self.f_desc.set(desc); self.f_valor.set(f'{val:.2f}'.replace('.',',')); self.editing_id=i
        self.entry_btn.configure(text='✓  Salvar alterações'); self.form_note.configure(text='Editando um lançamento: altere os campos acima e salve (ou clique em «Limpar formulário» para cancelar).')

    def delete_entry(self):
        sel=self.entries.selection()
        if not sel: messagebox.showinfo('Excluir lançamentos','Selecione um ou mais lançamentos na tabela. Dica: segure Ctrl ou Shift para marcar vários.'); return
        if messagebox.askyesno('Confirmar exclusão',f'Excluir {len(sel)} lançamento(s) selecionado(s)?'):
            ids=[int(i) for i in sel]
            for i in ids: self.con.execute('DELETE FROM lancamentos WHERE id=?',(i,))
            self.con.commit()
            if self.editing_id in ids: self.reset_form()
            self.refresh_page()

    # ------------------------------------------------------------ limpar e recomeçar
    def scope_filter(self,scope):
        month=self.get_month()
        if scope=='month': return "substr(data,1,7)=?",(month,),f'todos os lançamentos de {month_label(month)}'
        if scope=='year': return "substr(data,1,4)=?",(month[:4],),f'todos os lançamentos de {month[:4]} e a sobra inicial informada para esse ano'
        if scope=='imported': return "origem='planilha'",(),'todos os valores importados da planilha Excel (os lançamentos manuais são mantidos)'
        raise ValueError(scope)

    def clear_scope(self,scope):
        where,args,_=self.scope_filter(scope)
        n=self.con.execute(f'SELECT COUNT(*) FROM lancamentos WHERE {where}',args).fetchone()[0]
        self.con.execute(f'DELETE FROM lancamentos WHERE {where}',args)
        if scope=='year': self.con.execute('DELETE FROM saldos_iniciais WHERE substr(mes,1,4)=?',args)
        self.con.commit(); return n

    def clear_entries(self,scope):
        where,args,label=self.scope_filter(scope)
        n=self.con.execute(f'SELECT COUNT(*) FROM lancamentos WHERE {where}',args).fetchone()[0]
        if n==0: messagebox.showinfo('Nada para limpar','Não há lançamentos nesse período.'); return
        extra='\n\nUma cópia de segurança será criada automaticamente antes.' if scope in ('year','imported') else ''
        if not messagebox.askyesno('Confirmar limpeza',f'Vou apagar {label} ({n} registro(s)). Esta ação não pode ser desfeita.{extra}\n\nContinuar?',icon='warning',default='no'): return
        if scope in ('year','imported'): self.safety_backup(scope)
        n=self.clear_scope(scope); self.editing_id=None; self.refresh_page()
        messagebox.showinfo('Limpeza concluída',f'{n} lançamento(s) apagado(s).')

    def safety_backup(self,tag):
        """Cópia automática do banco antes de apagar em lote. Mantém só as 10 mais recentes."""
        self.con.commit(); path=DB.parent/f'backup_automatico_{tag}_{datetime.now():%Y%m%d_%H%M%S}.db'; shutil.copy2(DB,path)
        for old in sorted(DB.parent.glob('backup_automatico_*.db'),key=lambda p:p.stat().st_mtime)[:-10]:
            try: old.unlink()
            except OSError: pass
        return path

    def clear_goals(self):
        n=self.con.execute('SELECT COUNT(*) FROM metas').fetchone()[0]
        if n==0: messagebox.showinfo('Nada para limpar','Não há metas cadastradas.'); return
        if not messagebox.askyesno('Limpar metas',f'Apagar todas as {n} meta(s)? Esta ação não pode ser desfeita.',icon='warning',default='no'): return
        self.con.execute('DELETE FROM metas'); self.con.commit(); self.refresh_page()

    def restore_targets(self):
        if not messagebox.askyesno('Restaurar percentuais','Voltar os percentuais do orçamento para o padrão 55 / 20 / 10 / 10 / 5?'): return
        for k,v in TARGETS_DEFAULT.items(): self.con.execute('INSERT OR REPLACE INTO preferencias(chave,valor) VALUES(?,?)',(f'alvo_{k}',str(v)))
        self.con.commit(); self.show_page(self.current_page)

    def wipe_all(self):
        for t in ('lancamentos','metas','preferencias','saldos_iniciais'): self.con.execute(f'DELETE FROM {t}')
        for k,v in TARGETS_DEFAULT.items(): self.con.execute('INSERT OR REPLACE INTO preferencias(chave,valor) VALUES(?,?)',(f'alvo_{k}',str(v)))
        self.con.commit()

    def reset_all(self):
        if not messagebox.askyesno('Restaurar o programa do zero','Isto apaga TODOS os lançamentos, metas, sobra inicial e percentuais e deixa o programa como novo.\n\nUma cópia de segurança será criada automaticamente antes.\n\nDeseja continuar?',icon='warning',default='no'): return
        word=simpledialog.askstring('Confirmação final','Para confirmar, digite APAGAR TUDO (em maiúsculas):',parent=self)
        if (word or '').strip()!='APAGAR TUDO': messagebox.showinfo('Restauração cancelada','Nada foi apagado.'); return
        try: path=self.safety_backup('antes_de_restaurar')
        except OSError as e: messagebox.showerror('Não foi possível criar a cópia de segurança',f'Por segurança, nada foi apagado.\n\n{e}'); return
        self.wipe_all(); self.editing_id=None; self.top_month.set(date.today().strftime('%Y-%m')); self.show_page('Painel')
        messagebox.showinfo('Programa restaurado',f'O programa voltou ao estado inicial.\n\nCópia dos dados anteriores:\n{path}\n\nPara recuperá-los, use «Relatórios → Restaurar backup» e escolha esse arquivo.')

    def delete_entry(self):
        sel=self.entries.selection()
        if sel and messagebox.askyesno('Confirmar exclusão','Excluir o lançamento selecionado?'):
            self.con.execute('DELETE FROM lancamentos WHERE id=?',(int(sel[0]),)); self.con.commit(); self.refresh_page()

    def import_workbook(self):
        path=filedialog.askopenfilename(title='Selecione a planilha de organização financeira',filetypes=[('Planilhas Excel','*.xlsx *.xlsm')])
        if not path: return
        try:
            year,entries,opening,targets=extract_workbook_data(path)
            if not messagebox.askyesno('Importar planilha',f'Vou importar os lançamentos preenchidos da planilha:\n{Path(path).name}\n\nRegistros importados anteriormente (e os ajustes feitos na tela Planilha para o mesmo ano) serão substituídos. Os lançamentos manuais serão preservados. Continuar?'): return
            con=self.con; con.execute('BEGIN')
            con.execute("DELETE FROM lancamentos WHERE origem='planilha' OR (origem='grade' AND substr(data,1,4)=?)",(str(year),))
            for data,tipo,categoria,valor,grupo in entries:
                con.execute('INSERT INTO lancamentos(data,tipo,categoria,descricao,valor,origem,grupo) VALUES(?,?,?,?,?,?,?)',(data,tipo,categoria,'Importado da planilha',valor,'planilha',grupo))
            if opening:
                con.execute('INSERT OR REPLACE INTO saldos_iniciais(mes,valor) VALUES(?,?)',(f'{year}-01',opening))
            for key,value in targets.items():
                con.execute('INSERT OR REPLACE INTO preferencias(chave,valor) VALUES(?,?)',(f'alvo_{key}',str(value)))
            con.commit(); self.top_month.set(f'{year}-{date.today().month:02d}'); self.show_page(self.current_page)
            messagebox.showinfo('Importação concluída',f'Importação concluída.\n{len(entries)} lançamentos importados.\nAno da planilha: {year}.\n\nLinhas de totais e fórmulas foram ignoradas para evitar duplicidade.')
        except Exception as e:
            try: self.con.rollback()
            except Exception: pass
            messagebox.showerror('Erro na importação',f'Não foi possível importar a planilha.\n\nDetalhes: {e}')

    def export_excel(self):
        path=filedialog.asksaveasfilename(title='Salvar relatório Excel',defaultextension='.xlsx',filetypes=[('Excel','*.xlsx')],initialfile=f'Relatorio_Financeiro_{self.get_month()}.xlsx')
        if not path: return
        try:
            cur='"R$" #,##0.00'; month=self.get_month()
            wb=Workbook(); ws=wb.active; ws.title='Lançamentos'; ws.append(['Data','Tipo','Categoria','Descrição','Valor (R$)','Origem'])
            for r in self.con.execute('SELECT data,tipo,categoria,descricao,valor,origem FROM lancamentos ORDER BY data,id'): ws.append(list(r))
            for row in ws.iter_rows(min_row=2,min_col=5,max_col=5):
                for c in row: c.number_format=cur
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            ws2=wb.create_sheet('Metas'); ws2.append(['Meta','Valor alvo','Valor guardado','Falta','Prazo'])
            for nome,alvo,atual,prazo in self.con.execute('SELECT nome,valor_alvo,valor_atual,prazo FROM metas'): ws2.append([nome,alvo,atual,max(0,alvo-atual),prazo])
            for row in ws2.iter_rows(min_row=2,min_col=2,max_col=4):
                for c in row: c.number_format=cur
            ws3=wb.create_sheet('Resumo mensal'); ws3.append(['Mês','Receitas','Despesas','Investimentos','Saldo do mês'])
            for y,m in self.con.execute("SELECT DISTINCT substr(data,1,4),substr(data,6,2) FROM lancamentos ORDER BY 1,2"):
                mm=f'{y}-{m}'; s=self.sums(mm); ws3.append([mm,s['Receita'],s['Despesa'],s['Investimento'],s['Receita']-s['Despesa']-s['Investimento']])
            for row in ws3.iter_rows(min_row=2,min_col=2,max_col=5):
                for c in row: c.number_format=cur
            n=ws3.max_row
            if n>1:
                bar=XlBarChart(); bar.type='col'; bar.title='Receitas, despesas e investimentos por mês'; bar.height=9; bar.width=20
                bar.add_data(Reference(ws3,min_col=2,max_col=4,min_row=1,max_row=n),titles_from_data=True); bar.set_categories(Reference(ws3,min_col=1,min_row=2,max_row=n)); ws3.add_chart(bar,'G2')
            ws4=wb.create_sheet('Blocos do mês'); ws4.append(['Bloco',f'Gasto em {month} (R$)'])
            for g,v in self.group_spend(month).items(): ws4.append([g,v])
            for row in ws4.iter_rows(min_row=2,min_col=2,max_col=2):
                for c in row: c.number_format=cur
            if any(v>0 for v in self.group_spend(month).values()):
                pie=XlPieChart(); pie.title=f'Gastos por bloco · {month}'; pie.height=9; pie.width=14
                pie.add_data(Reference(ws4,min_col=2,min_row=1,max_row=ws4.max_row),titles_from_data=True); pie.set_categories(Reference(ws4,min_col=1,min_row=2,max_row=ws4.max_row)); ws4.add_chart(pie,'D2')
            for sh in wb.worksheets:
                for c in sh[1]:
                    f=copy(c.font); f.bold=True; c.font=f
                for col in sh.columns:
                    letter=col[0].column_letter; sh.column_dimensions[letter].width=min(38,max(14,max(len(str(c.value or '')) for c in col)+2))
            wb.save(path); messagebox.showinfo('Exportação concluída',f'Relatório salvo em:\n{path}')
        except Exception as e: messagebox.showerror('Erro ao exportar',str(e))

    def export_csv(self):
        path=filedialog.asksaveasfilename(title='Salvar lançamentos CSV',defaultextension='.csv',filetypes=[('CSV','*.csv')],initialfile='Lancamentos_Financeiros.csv')
        if not path: return
        with open(path,'w',newline='',encoding='utf-8-sig') as f:
            writer=csv.writer(f,delimiter=';'); writer.writerow(['Data','Tipo','Categoria','Descrição','Valor (R$)','Origem']); writer.writerows(self.con.execute('SELECT data,tipo,categoria,descricao,valor,origem FROM lancamentos ORDER BY data,id'))
        messagebox.showinfo('Exportação concluída',f'Arquivo CSV salvo em:\n{path}')

    def backup_menu(self):
        self.show_page('Relatórios')

    def backup_db(self):
        path=filedialog.asksaveasfilename(title='Salvar backup do banco',defaultextension='.db',filetypes=[('Banco SQLite','*.db')],initialfile=f'Backup_Financas_{date.today().isoformat()}.db')
        if not path: return
        self.con.commit(); shutil.copy2(DB,path); messagebox.showinfo('Backup criado',f'Cópia salva em:\n{path}')

    def restore_db(self):
        path=filedialog.askopenfilename(title='Selecionar backup do banco',filetypes=[('Banco SQLite','*.db')])
        if not path: return
        if not messagebox.askyesno('Restaurar backup','A restauração substituirá os dados atuais. Faça backup antes. Continuar?'): return
        try:
            self.con.close(); shutil.copy2(path,DB); self.con=connect(); self.show_page(self.current_page); messagebox.showinfo('Restauração concluída','Backup restaurado com sucesso.')
        except Exception as e:
            self.con=connect(); messagebox.showerror('Erro na restauração',str(e))

    def close(self):
        self.con.close(); self.destroy()

if __name__=='__main__':
    App().mainloop()
